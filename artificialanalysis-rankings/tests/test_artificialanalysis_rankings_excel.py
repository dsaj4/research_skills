from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import generate_artificialanalysis_rankings_excel as rankings


def sample_rows() -> list[dict[str, object]]:
    generated_at = "2026-04-02 10:00:00 CST"
    rows = []
    for idx in range(12):
        rows.append(
            {
                "rank": idx + 1,
                "model_id": f"model-{idx + 1}",
                "name": f"Model {idx + 1}",
                "organization": "Example Org",
                "organization_id": "example-org",
                "primary_score": 100 - idx,
                "secondary_score": 90 - idx,
                "benchmark_one_score": 0.90 - (idx * 0.01),
                "benchmark_two_score": 0.80 - (idx * 0.01),
                "reference_score": 0.50 - (idx * 0.01),
                "context": 1000000 - (idx * 10000),
                "context_display": rankings.format_context(1000000 - (idx * 10000)),
                "input_price": 1.0 + idx,
                "input_price_display": rankings.format_price(1.0 + idx),
                "output_price": 2.0 + idx,
                "output_price_display": rankings.format_price(2.0 + idx),
                "license": "Proprietary",
                "announcement_date": "2026-04-01",
                "source_url": rankings.SOURCE_URL,
                "generated_at": generated_at,
            }
        )
    return rows


class GeneratedProjectTests(unittest.TestCase):
    def test_constants_match_scaffolded_site(self) -> None:
        self.assertEqual(rankings.SOURCE_URL, "https://artificialanalysis.ai/")
        self.assertEqual(rankings.DISPLAY_PAGE_TITLE, "Artificial Analysis Main Leaderboards")

    def test_extract_source_rows_from_html_merges_homepage_datasets(self) -> None:
        html = """
        <script type="application/ld+json">
        {"@type":"Dataset","name":"Intelligence","data":[
            {"modelName":"Model A","intelligenceIndex":57.1,"detailsUrl":"/models/model-a"},
            {"modelName":"Model B","intelligenceIndex":42.0,"detailsUrl":"/models/model-b"}
        ]}
        </script>
        <script type="application/ld+json">
        {"@type":"Dataset","name":"Speed","data":[
            {"modelName":"Model A","medianOutputSpeed":123.4,"detailsUrl":"/models/model-a"},
            {"modelName":"Model B","medianOutputSpeed":99.9,"detailsUrl":"/models/model-b"}
        ]}
        </script>
        <script type="application/ld+json">
        {"@type":"Dataset","name":"Price","data":[
            {"modelName":"Model A","pricePerMillionTokens":4.5,"detailsUrl":"/models/model-a"},
            {"modelName":"Model B","pricePerMillionTokens":1.5,"detailsUrl":"/models/model-b"}
        ]}
        </script>
        """

        rows = rankings.extract_source_rows_from_html(html)
        normalized = rankings.normalize_source_rows(rows, generated_at="2026-04-02 10:00:00 CST")

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["model_id"], "model-a")
        self.assertEqual(rows[0]["intelligence_index"], 57.1)
        self.assertEqual(rows[0]["output_speed"], 123.4)
        self.assertEqual(rows[0]["price_per_million_tokens"], 4.5)
        self.assertEqual(normalized[0]["primary_score"], 57.1)
        self.assertEqual(normalized[0]["secondary_score"], 123.4)
        self.assertEqual(normalized[0]["input_price"], 4.5)
        self.assertEqual(normalized[0]["context"], 57.1)

    def test_build_workbook_creates_display(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "artificialanalysis_rankings_models.xlsx"
            rankings.build_workbook(sample_rows(), workbook_path, generated_at="2026-04-02 10:00:00 CST", include_display=True)
            wb = load_workbook(workbook_path)
            ws = wb[rankings.DISPLAY_SHEET_TITLE]
            self.assertEqual(ws["A1"].value, rankings.DISPLAY_PAGE_TITLE)
            self.assertEqual(ws["B4"].value, rankings.DISPLAY_TOP_LEFT_TITLE)
            self.assertEqual(len(ws._charts), 4)
            self.assertEqual(wb[rankings.RAW_MODELS_SHEET]["A2"].value, 1)
            self.assertEqual(wb[rankings.DATA_PRIMARY_SHEET]["A2"].value, "Model 1")
            self.assertEqual(wb[rankings.DATA_BENCHMARKS_SHEET]["A2"].value, "Model 12")


if __name__ == "__main__":
    unittest.main()
