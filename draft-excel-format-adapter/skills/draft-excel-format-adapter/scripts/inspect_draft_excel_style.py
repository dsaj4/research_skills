#!/usr/bin/env python
import argparse
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


def workbook_summary(path):
    wb = load_workbook(path, read_only=False, data_only=False)
    lines = [f"workbook: {path}", f"sheets: {wb.sheetnames}"]
    for ws in wb.worksheets:
        lines.append(
            f"- {ws.title}: rows={ws.max_row} cols={ws.max_column} "
            f"charts={len(getattr(ws, '_charts', []))} images={len(getattr(ws, '_images', []))} "
            f"merged={len(ws.merged_cells.ranges)}"
        )
    return lines


def zip_colors(path):
    colors = set()
    fills = set()
    fonts = set()
    with zipfile.ZipFile(path) as zf:
        for name in zf.namelist():
            if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                xml = zf.read(name).decode("utf-8", errors="ignore")
                colors.update(re.findall(r'<a:srgbClr val="([A-Fa-f0-9]+)"', xml))
        if "xl/styles.xml" in zf.namelist():
            styles = zf.read("xl/styles.xml").decode("utf-8", errors="ignore")
            fills.update(re.findall(r'rgb="([A-F0-9]{8})"', styles))
            fonts.update(re.findall(r"<name val=\"([^\"]+)\"", styles))
    return sorted(colors), sorted(fills), sorted(fonts)


def main():
    parser = argparse.ArgumentParser(description="Inspect draft/template Excel workbook style signals.")
    parser.add_argument("--workbook", required=True, type=Path, help="Workbook to inspect.")
    args = parser.parse_args()

    for line in workbook_summary(args.workbook):
        print(line)
    chart_colors, fill_colors, fonts = zip_colors(args.workbook)
    print(f"chart_colors: {chart_colors}")
    print(f"fill_colors: {fill_colors}")
    print(f"fonts: {fonts}")


if __name__ == "__main__":
    main()
