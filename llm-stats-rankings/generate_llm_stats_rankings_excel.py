from __future__ import annotations

import argparse
import csv
import json
import subprocess
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_URL = "https://llm-stats.com/"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/146.0.0.0 Safari/537.36"
)

DISPLAY_SHEET_TITLE = "Display"
META_SHEET = "Meta"
RAW_MODELS_SHEET = "Raw_Models"
DATA_TOP20_SHEET = "Data_Top20"
DATA_CODING_SHEET = "Data_Coding"
DATA_CHAT_SHEET = "Data_Chat"
DATA_BENCHMARKS_SHEET = "Data_Benchmarks"
DATA_PRICE_CONTEXT_SHEET = "Data_PriceContext"

DISPLAY_PAGE_TITLE = "LLM Stats Dashboard"
DISPLAY_PAGE_SUBTITLE = (
    "Display layer mirrors openrouter-rankings; only helper-sheet datasets should change."
)
DISPLAY_TOP_LEFT_TITLE = "[Coding Arena] Top 10"
DISPLAY_TOP_RIGHT_TITLE = "[Chat Arena] Top 10"
DISPLAY_BOTTOM_LEFT_TITLE = "[Benchmarks] GPQA / SWE-Bench"
DISPLAY_BOTTOM_RIGHT_TITLE = "[Price vs Context] Cost-Context Snapshot"
DISPLAY_REFRESHED_AT_LABEL = "Display refreshed: "
DISPLAY_DATA_REFRESHED_LABEL = "Data refreshed: "
DISPLAY_PLACEHOLDER_NOTE = "Data layer refreshed. Run chart-refresh to rebuild the display page."
DISPLAY_MODEL_COUNT_LABEL = "Model count: "
DISPLAY_CHART_SCOPE_LABEL = "Chart scope: "
DISPLAY_BENCHMARK_SUMMARY_TITLE = "Benchmark summary"
DISPLAY_PRICE_CONTEXT_SUMMARY_TITLE = "Price / Context summary"
DISPLAY_TOP_SUMMARY_ROWS = 8

RAW_MODEL_HEADERS = [
    "rank",
    "model_id",
    "name",
    "organization",
    "organization_id",
    "chat_arena_score",
    "coding_arena_score",
    "gpqa_score",
    "swe_bench_verified_score",
    "hle_score",
    "context",
    "context_display",
    "input_price",
    "input_price_display",
    "output_price",
    "output_price_display",
    "license",
    "announcement_date",
    "source_url",
    "generated_at",
]

SUMMARY_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
TABLE_HEADER_FILL = PatternFill("solid", fgColor="1E293B")
SUMMARY_HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")
SUMMARY_NOTE_FILL = PatternFill("solid", fgColor="F8FAFC")
OPENROUTER_HIGHLIGHT_BLUE = "196BA5"
OPENROUTER_LEADERBOARD_BLUE = "7699C8"
OPENROUTER_ACCENT_GOLD = "948A54"
OPENROUTER_MUTED_GRAY = "B6B6B6"
OPENROUTER_LIGHT_BLUE = "B9CDE5"
SERIES_FALLBACK_COLORS = [
    "0088FE",
    "00C49F",
    "FFBB28",
    "FF8042",
    "FF6347",
    "4682B4",
    "9ACD32",
    "DA70D6",
    "40E0D0",
    "FF69B4",
]


def current_timestamp() -> str:
    return datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M:%S %Z")


def fetch_homepage_html(url: str = SOURCE_URL) -> str:
    command = [
        "curl.exe",
        "-L",
        url,
        "-H",
        f"User-Agent: {USER_AGENT}",
        "-H",
        f"Referer: {SOURCE_URL}",
        "--max-time",
        "180",
        "--silent",
        "--show-error",
    ]
    result = subprocess.run(command, check=True, capture_output=True)
    return result.stdout.decode("utf-8", errors="replace")


def load_homepage_html(cache_path: Path) -> str:
    try:
        html = fetch_homepage_html()
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(html, encoding="utf-8")
        return html
    except Exception:
        if cache_path.exists():
            return cache_path.read_text(encoding="utf-8")
        raise


def extract_balanced_block(
    text: str,
    start_idx: int,
    open_char: str = "[",
    close_char: str = "]",
) -> str:
    depth = 0
    in_string = False
    escaped = False
    for idx in range(start_idx, len(text)):
        char = text[idx]
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if char == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if char == open_char:
            depth += 1
        elif char == close_char:
            depth -= 1
            if depth == 0:
                return text[start_idx : idx + 1]
    raise ValueError("Unbalanced block while extracting homepage payload")


def extract_initial_homepage_models_from_html(html: str) -> list[dict[str, object]]:
    marker = 'initialHomepageLLMModels\\":'
    marker_idx = html.find(marker)
    if marker_idx == -1:
        raise ValueError("Could not locate initialHomepageLLMModels in homepage HTML")

    array_start = html.find("[", marker_idx)
    if array_start == -1:
        raise ValueError("Could not locate homepage models array start")

    raw_array = extract_balanced_block(html, array_start, "[", "]")
    decoded_array = bytes(raw_array, "utf-8").decode("unicode_escape")
    rows = json.loads(decoded_array)
    if not isinstance(rows, list):
        raise ValueError("Homepage payload is not a list")
    return rows


def format_context(value: int | float | None) -> str:
    if value in (None, ""):
        return ""
    numeric = float(value)
    if numeric >= 1_000_000:
        return f"{numeric / 1_000_000:.1f}M"
    if numeric >= 1_000:
        return f"{numeric / 1_000:.0f}K"
    return f"{numeric:.0f}"


def format_price(value: int | float | None) -> str:
    if value in (None, ""):
        return ""
    return f"${float(value):.2f}"


def normalize_homepage_models(
    rows: list[dict[str, object]],
    generated_at: str,
) -> list[dict[str, object]]:
    normalized: list[dict[str, object]] = []
    for idx, row in enumerate(rows):
        arena_scores = row.get("arena_scores") or {}
        context = row.get("context")
        input_price = row.get("input_price")
        output_price = row.get("output_price")
        normalized.append(
            {
                "rank": idx + 1,
                "model_id": row.get("model_id"),
                "name": row.get("name"),
                "organization": row.get("organization"),
                "organization_id": row.get("organization_id"),
                "chat_arena_score": arena_scores.get("chat-arena"),
                "coding_arena_score": arena_scores.get("coding-arena"),
                "gpqa_score": row.get("gpqa_score"),
                "swe_bench_verified_score": row.get("swe_bench_verified_score"),
                "hle_score": row.get("hle_score"),
                "context": context,
                "context_display": format_context(context),
                "input_price": input_price,
                "input_price_display": format_price(input_price),
                "output_price": output_price,
                "output_price_display": format_price(output_price),
                "license": "Open Source" if row.get("is_open_source") else "Proprietary",
                "announcement_date": row.get("announcement_date"),
                "source_url": SOURCE_URL,
                "generated_at": generated_at,
            }
        )
    return normalized


def write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=RAW_MODEL_HEADERS)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in RAW_MODEL_HEADERS})


def safe_float(value: object) -> float | None:
    if value in (None, ""):
        return None
    return float(value)


def format_score(value: object, decimals: int = 2) -> str:
    numeric = safe_float(value)
    if numeric is None:
        return "-"
    return f"{numeric:.{decimals}f}"


def format_ratio_score(value: object) -> str:
    numeric = safe_float(value)
    if numeric is None:
        return "-"
    return f"{numeric:.3f}"


def clear_sheet(worksheet) -> None:
    for merged_range in list(worksheet.merged_cells.ranges):
        worksheet.unmerge_cells(str(merged_range))
    if worksheet.max_row:
        worksheet.delete_rows(1, worksheet.max_row)
    worksheet._charts = []
    worksheet.freeze_panes = None


def create_workbook_skeleton() -> Workbook:
    workbook = Workbook()
    workbook.active.title = DISPLAY_SHEET_TITLE
    workbook.create_sheet(META_SHEET)
    workbook.create_sheet(RAW_MODELS_SHEET)
    workbook.create_sheet(DATA_TOP20_SHEET)
    workbook.create_sheet(DATA_CODING_SHEET)
    workbook.create_sheet(DATA_CHAT_SHEET)
    workbook.create_sheet(DATA_BENCHMARKS_SHEET)
    workbook.create_sheet(DATA_PRICE_CONTEXT_SHEET)
    workbook.active = workbook.sheetnames.index(DISPLAY_SHEET_TITLE)
    return workbook


def ensure_workbook_skeleton(workbook: Workbook) -> Workbook:
    required_titles = [
        DISPLAY_SHEET_TITLE,
        META_SHEET,
        RAW_MODELS_SHEET,
        DATA_TOP20_SHEET,
        DATA_CODING_SHEET,
        DATA_CHAT_SHEET,
        DATA_BENCHMARKS_SHEET,
        DATA_PRICE_CONTEXT_SHEET,
    ]
    for title in required_titles:
        if title not in workbook.sheetnames:
            workbook.create_sheet(title)
    if "Charts" in workbook.sheetnames:
        workbook.remove(workbook["Charts"])
    workbook.active = workbook.sheetnames.index(DISPLAY_SHEET_TITLE)
    return workbook


def read_meta_values(workbook: Workbook) -> dict[str, str]:
    if META_SHEET not in workbook.sheetnames:
        return {}
    worksheet = workbook[META_SHEET]
    values: dict[str, str] = {}
    for key, value in worksheet.iter_rows(min_row=2, max_col=2, values_only=True):
        if key:
            values[str(key)] = "" if value is None else str(value)
    return values


def load_rows_from_worksheet(worksheet) -> list[dict[str, object]]:
    headers = [cell.value for cell in worksheet[1]]
    if not any(headers):
        return []

    rows: list[dict[str, object]] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        rows.append(
            {
                str(headers[idx]): value
                for idx, value in enumerate(values)
                if idx < len(headers) and headers[idx] is not None
            }
        )
    return rows


def load_models_from_workbook(workbook_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(workbook_path)
    return load_rows_from_worksheet(workbook[RAW_MODELS_SHEET])


def write_table_sheet(
    worksheet,
    headers: list[str],
    rows: list[list[object]],
    number_formats: dict[int, str] | None = None,
) -> None:
    clear_sheet(worksheet)
    text_light = "F8FAFC"

    for col_idx, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color=text_light)
        cell.fill = TABLE_HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="center")
            if number_formats and col_idx in number_formats:
                cell.number_format = number_formats[col_idx]

    worksheet.freeze_panes = "A2"
    for col_idx in range(1, len(headers) + 1):
        worksheet.column_dimensions[get_column_letter(col_idx)].width = 18


def write_raw_models_sheet(worksheet, rows: list[dict[str, object]]) -> None:
    raw_rows = [[row.get(header) for header in RAW_MODEL_HEADERS] for row in rows]
    write_table_sheet(
        worksheet,
        RAW_MODEL_HEADERS,
        raw_rows,
        number_formats={
            6: "0.00",
            7: "0.00",
            8: "0.000",
            9: "0.000",
            10: "0.000",
            13: "$0.00",
            15: "$0.00",
        },
    )
    widths = {
        "A": 8,
        "B": 24,
        "C": 28,
        "D": 18,
        "E": 18,
        "F": 16,
        "G": 16,
        "H": 14,
        "I": 18,
        "J": 12,
        "K": 12,
        "L": 14,
        "M": 12,
        "N": 12,
        "O": 12,
        "P": 12,
        "Q": 14,
        "R": 16,
        "S": 28,
        "T": 24,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def write_meta_sheet(
    worksheet,
    generated_at: str,
    row_count: int,
    data_refreshed_at: str,
    display_refreshed_at: str,
) -> None:
    clear_sheet(worksheet)
    text_light = "F8FAFC"

    worksheet["A1"] = "LLM Stats Snapshot"
    worksheet["A1"].font = Font(size=16, bold=True, color=text_light)
    worksheet["A1"].fill = SUMMARY_FILL

    meta_rows = [
        ("source_url", SOURCE_URL),
        ("generated_at", generated_at),
        ("data_refreshed_at", data_refreshed_at),
        ("display_refreshed_at", display_refreshed_at),
        ("row_count", row_count),
        ("extraction_method", "Embedded initialHomepageLLMModels payload parsed from homepage HTML."),
        ("cache_fallback", "Use tmp/api-cache/llm_stats_homepage.html if live fetch fails."),
        ("chart_refresh_mode", "chart-refresh rebuilds helper sheets and display layout from Raw_Models."),
        ("notes", "Single workbook with raw model rows plus an openrouter-style dashboard layer."),
    ]
    for row_idx, (label, value) in enumerate(meta_rows, start=2):
        label_cell = worksheet.cell(row=row_idx, column=1, value=label)
        value_cell = worksheet.cell(row=row_idx, column=2, value=value)
        label_cell.font = Font(bold=True)
        label_cell.fill = HEADER_FILL
        value_cell.fill = HEADER_FILL
        value_cell.alignment = Alignment(wrap_text=True)

    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 96


def top_n_by_metric(
    rows: list[dict[str, object]],
    metric: str,
    limit: int = 20,
) -> list[dict[str, object]]:
    filtered = [row for row in rows if row.get(metric) is not None]
    filtered.sort(key=lambda row: float(row[metric]), reverse=True)
    return filtered[:limit]


def top_n_by_context(rows: list[dict[str, object]], limit: int = 20) -> list[dict[str, object]]:
    filtered = [
        row
        for row in rows
        if row.get("context") is not None and row.get("input_price") is not None
    ]
    filtered.sort(
        key=lambda row: (float(row["context"]), -float(row["input_price"])),
        reverse=True,
    )
    return filtered[:limit]


def top_n_by_benchmark_blend(
    rows: list[dict[str, object]],
    limit: int = 20,
) -> list[dict[str, object]]:
    filtered = [
        row
        for row in rows
        if row.get("gpqa_score") is not None or row.get("swe_bench_verified_score") is not None
    ]

    def blended_score(row: dict[str, object]) -> float:
        values = [
            float(row[key])
            for key in ("gpqa_score", "swe_bench_verified_score")
            if row.get(key) is not None
        ]
        return sum(values) / len(values)

    filtered.sort(key=blended_score, reverse=True)
    return filtered[:limit]


def row_color_map(rows: list[dict[str, object]]) -> dict[str, str]:
    return {
        str(row.get("name", f"Series {idx + 1}")): SERIES_FALLBACK_COLORS[idx % len(SERIES_FALLBACK_COLORS)]
        for idx, row in enumerate(rows)
    }


def best_row_by_metric(rows: list[dict[str, object]], metric: str) -> dict[str, object] | None:
    candidates = [row for row in rows if row.get(metric) is not None]
    if not candidates:
        return None
    return max(candidates, key=lambda row: float(row[metric]))


def format_model_metric(
    row: dict[str, object] | None,
    metric: str,
    formatter,
) -> str:
    if not row:
        return "-"
    return f"{row.get('name')}, {formatter(row.get(metric))}"


def prepare_display_datasets(rows: list[dict[str, object]]) -> dict[str, object]:
    top20_rows = rows[:20]
    coding_rows = top_n_by_metric(rows, "coding_arena_score", limit=10)
    chat_rows = top_n_by_metric(rows, "chat_arena_score", limit=10)
    benchmark_rows = top_n_by_benchmark_blend(top20_rows, limit=10)
    price_context_rows = top_n_by_context(top20_rows, limit=10)

    coding_summary = [
        {
            "rank": idx,
            "model": row["name"],
            "value": format_score(row.get("coding_arena_score")),
        }
        for idx, row in enumerate(coding_rows[:DISPLAY_TOP_SUMMARY_ROWS], start=1)
    ]
    chat_summary = [
        {
            "rank": idx,
            "model": row["name"],
            "value": format_score(row.get("chat_arena_score")),
        }
        for idx, row in enumerate(chat_rows[:DISPLAY_TOP_SUMMARY_ROWS], start=1)
    ]
    benchmark_summary = [
        {
            "label": "Best GPQA",
            "value": format_model_metric(best_row_by_metric(rows, "gpqa_score"), "gpqa_score", format_ratio_score),
        },
        {
            "label": "Best SWE-Bench",
            "value": format_model_metric(
                best_row_by_metric(rows, "swe_bench_verified_score"),
                "swe_bench_verified_score",
                format_ratio_score,
            ),
        },
        {
            "label": "Best HLE",
            "value": format_model_metric(best_row_by_metric(rows, "hle_score"), "hle_score", format_ratio_score),
        },
        {
            "label": "Benchmark models",
            "value": str(len(benchmark_rows)),
        },
    ]
    price_context_summary = [
        {
            "rank": idx,
            "model": row["name"],
            "value": f"{row.get('context_display') or format_context(row.get('context'))} / {format_price(row.get('input_price'))}",
        }
        for idx, row in enumerate(price_context_rows[:8], start=1)
    ]

    return {
        "top20": top20_rows,
        "coding": coding_rows,
        "chat": chat_rows,
        "benchmarks": benchmark_rows,
        "price_context": price_context_rows,
        "coding_summary": coding_summary,
        "chat_summary": chat_summary,
        "benchmark_summary": benchmark_summary,
        "price_context_summary": price_context_summary,
        "model_count": len(rows),
        "chart_scope": "Top 10 models per panel",
    }


def write_ranked_metric_sheet(
    worksheet,
    rows: list[dict[str, object]],
    metric_key: str,
    metric_header: str,
    number_format: str = "0.00",
) -> None:
    headers = ["model", metric_header]
    values = [[row.get("name"), row.get(metric_key)] for row in rows]
    write_table_sheet(
        worksheet,
        headers,
        values,
        number_formats={2: number_format},
    )
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 16


def write_benchmark_chart_sheet(worksheet, rows: list[dict[str, object]]) -> None:
    headers = ["model", "GPQA", "SWE-Bench"]
    values = [
        [row.get("name"), row.get("gpqa_score"), row.get("swe_bench_verified_score")]
        for row in rows
    ]
    write_table_sheet(
        worksheet,
        headers,
        values,
        number_formats={2: "0.000", 3: "0.000"},
    )
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 12
    worksheet.column_dimensions["C"].width = 14


def write_price_context_chart_sheet(worksheet, rows: list[dict[str, object]]) -> None:
    table_rows = [[row.get("name"), row.get("context"), row.get("input_price")] for row in rows]
    write_table_sheet(
        worksheet,
        ["name", "context", "input_price"],
        table_rows,
        number_formats={2: "#,##0", 3: "$0.00"},
    )
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 14
    worksheet.column_dimensions["C"].width = 14


def write_display_helper_sheets(workbook: Workbook, display_data: dict[str, object]) -> None:
    top20_rows = [
        [
            row.get("rank"),
            row.get("name"),
            row.get("organization"),
            row.get("coding_arena_score"),
            row.get("chat_arena_score"),
            row.get("gpqa_score"),
            row.get("swe_bench_verified_score"),
        ]
        for row in display_data["top20"]
    ]
    write_table_sheet(
        workbook[DATA_TOP20_SHEET],
        [
            "rank",
            "name",
            "organization",
            "coding_arena_score",
            "chat_arena_score",
            "gpqa_score",
            "swe_bench_verified_score",
        ],
        top20_rows,
        number_formats={4: "0.00", 5: "0.00", 6: "0.000", 7: "0.000"},
    )
    write_ranked_metric_sheet(
        workbook[DATA_CODING_SHEET],
        display_data["coding"],
        "coding_arena_score",
        "coding_arena_score",
    )
    write_ranked_metric_sheet(
        workbook[DATA_CHAT_SHEET],
        display_data["chat"],
        "chat_arena_score",
        "chat_arena_score",
    )
    write_benchmark_chart_sheet(workbook[DATA_BENCHMARKS_SHEET], display_data["benchmarks"])
    write_price_context_chart_sheet(
        workbook[DATA_PRICE_CONTEXT_SHEET],
        display_data["price_context"],
    )
    for sheet_name in (
        DATA_TOP20_SHEET,
        DATA_CODING_SHEET,
        DATA_CHAT_SHEET,
        DATA_BENCHMARKS_SHEET,
        DATA_PRICE_CONTEXT_SHEET,
    ):
        workbook[sheet_name].sheet_state = "hidden"


def write_display_placeholder(worksheet, generated_at: str) -> None:
    clear_sheet(worksheet)
    worksheet.sheet_view.showGridLines = False
    worksheet["A1"] = DISPLAY_PAGE_TITLE
    worksheet["A1"].font = Font(size=16, bold=True)
    worksheet["A2"] = DISPLAY_PLACEHOLDER_NOTE
    worksheet["A3"] = f"{DISPLAY_DATA_REFRESHED_LABEL}{generated_at}"
    configure_display_page_setup(worksheet)


def write_display_summary_headers(
    worksheet,
    start_row: int,
    start_col: int,
    headers: list[str],
) -> None:
    for offset, header in enumerate(headers):
        cell = worksheet.cell(row=start_row, column=start_col + offset, value=header)
        cell.font = Font(bold=True, size=9)
        cell.fill = SUMMARY_HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def style_display_summary_range(
    worksheet,
    start_row: int,
    end_row: int,
    start_col: int,
    end_col: int,
) -> None:
    for row_idx in range(start_row, end_row + 1):
        worksheet.row_dimensions[row_idx].height = 15
        for col_idx in range(start_col, end_col + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical="center", shrink_to_fit=True)
            if cell.value is not None and col_idx != start_col:
                cell.fill = SUMMARY_NOTE_FILL


def apply_bar_series_colors(chart, colors: list[str]) -> None:
    for idx, series in enumerate(chart.series):
        color = colors[idx % len(colors)]
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color


def apply_scatter_series_style(chart, colors: list[str]) -> None:
    for idx, series in enumerate(chart.series):
        color = colors[idx % len(colors)]
        try:
            series.graphicalProperties.line.noFill = True
        except AttributeError:
            series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.solidFill = color
        try:
            series.marker.symbol = "circle"
            series.marker.size = 7
        except AttributeError:
            pass


def add_ranked_bar_chart(
    worksheet,
    data_sheet,
    anchor: str,
    series_color: str,
    axis_title: str,
) -> None:
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.width = 13.6
    chart.height = 6.0
    chart.legend = None
    chart.gapWidth = 35
    chart.x_axis.title = axis_title
    chart.y_axis.reverseOrder = True
    chart.add_data(
        Reference(data_sheet, min_col=2, min_row=1, max_row=data_sheet.max_row),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(data_sheet, min_col=1, min_row=2, max_row=data_sheet.max_row)
    )
    if chart.series:
        chart.series[0].graphicalProperties.solidFill = series_color
        chart.series[0].graphicalProperties.line.solidFill = series_color
    worksheet.add_chart(chart, anchor)


def add_benchmark_chart(worksheet, data_sheet, anchor: str) -> None:
    chart = BarChart()
    chart.type = "bar"
    chart.style = 11
    chart.width = 13.6
    chart.height = 6.0
    chart.legend.position = "r"
    chart.legend.overlay = False
    chart.gapWidth = 35
    chart.x_axis.title = "Benchmark score"
    chart.y_axis.reverseOrder = True
    chart.add_data(
        Reference(data_sheet, min_col=2, max_col=3, min_row=1, max_row=data_sheet.max_row),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(data_sheet, min_col=1, min_row=2, max_row=data_sheet.max_row)
    )
    benchmark_colors = [OPENROUTER_LEADERBOARD_BLUE, OPENROUTER_ACCENT_GOLD]
    apply_bar_series_colors(chart, benchmark_colors)
    worksheet.add_chart(chart, anchor)


def add_price_context_chart(worksheet, data_sheet, anchor: str) -> None:
    chart = ScatterChart()
    chart.width = 13.6
    chart.height = 6.0
    chart.scatterStyle = "marker"
    chart.x_axis.title = "Context"
    chart.y_axis.title = "Input $/M"
    chart.legend = None
    xvalues = Reference(data_sheet, min_col=2, min_row=2, max_row=data_sheet.max_row)
    yvalues = Reference(data_sheet, min_col=3, min_row=2, max_row=data_sheet.max_row)
    chart.series.append(Series(yvalues, xvalues, title="Models"))
    apply_scatter_series_style(chart, [OPENROUTER_HIGHLIGHT_BLUE])
    worksheet.add_chart(chart, anchor)


def build_display_sheet(worksheet, display_data: dict[str, object], generated_at: str) -> None:
    clear_sheet(worksheet)
    worksheet.sheet_view.showGridLines = False

    worksheet["A1"] = DISPLAY_PAGE_TITLE
    worksheet["A1"].font = Font(size=16, bold=True)
    worksheet["A2"] = DISPLAY_PAGE_SUBTITLE
    worksheet["A3"] = f"{DISPLAY_REFRESHED_AT_LABEL}{generated_at}"
    worksheet["B4"] = DISPLAY_TOP_LEFT_TITLE
    worksheet["M4"] = DISPLAY_TOP_RIGHT_TITLE
    worksheet["B30"] = DISPLAY_BOTTOM_LEFT_TITLE
    worksheet["M30"] = DISPLAY_BOTTOM_RIGHT_TITLE
    for cell_ref in ("B4", "M4", "B30", "M30"):
        worksheet[cell_ref].font = Font(bold=True, size=11)

    worksheet.column_dimensions["A"].width = 2
    worksheet.column_dimensions["B"].width = 9
    worksheet.column_dimensions["C"].width = 37
    worksheet.column_dimensions["D"].width = 18
    worksheet.column_dimensions["M"].width = 9
    worksheet.column_dimensions["N"].width = 37
    worksheet.column_dimensions["O"].width = 18
    for row_idx, height in {
        1: 24,
        2: 18,
        3: 16,
        4: 18,
        5: 8,
        19: 18,
        20: 16,
        30: 18,
        45: 18,
        46: 16,
    }.items():
        worksheet.row_dimensions[row_idx].height = height

    worksheet["B19"] = f"{DISPLAY_MODEL_COUNT_LABEL}{display_data['model_count']}"
    worksheet["M19"] = f"{DISPLAY_CHART_SCOPE_LABEL}{display_data['chart_scope']}"
    worksheet["B45"] = DISPLAY_BENCHMARK_SUMMARY_TITLE
    worksheet["M45"] = DISPLAY_PRICE_CONTEXT_SUMMARY_TITLE
    for cell_ref in ("B19", "M19", "B45", "M45"):
        worksheet[cell_ref].font = Font(bold=True, size=11)

    write_display_summary_headers(worksheet, 20, 2, ["Rank", "Model", "Coding"])
    for row_idx, summary in enumerate(display_data["coding_summary"], start=21):
        worksheet.cell(row=row_idx, column=2, value=summary["rank"])
        worksheet.cell(row=row_idx, column=3, value=summary["model"])
        worksheet.cell(row=row_idx, column=4, value=summary["value"])
    style_display_summary_range(worksheet, 21, 28, 2, 4)

    write_display_summary_headers(worksheet, 20, 13, ["Rank", "Model", "Chat"])
    for row_idx, summary in enumerate(display_data["chat_summary"], start=21):
        worksheet.cell(row=row_idx, column=13, value=summary["rank"])
        worksheet.cell(row=row_idx, column=14, value=summary["model"])
        worksheet.cell(row=row_idx, column=15, value=summary["value"])
    style_display_summary_range(worksheet, 21, 28, 13, 15)

    write_display_summary_headers(worksheet, 46, 2, ["Metric", "Value"])
    for row_idx, summary in enumerate(display_data["benchmark_summary"], start=47):
        worksheet.cell(row=row_idx, column=2, value=summary["label"])
        worksheet.cell(row=row_idx, column=3, value=summary["value"])
    style_display_summary_range(worksheet, 47, 50, 2, 3)

    write_display_summary_headers(worksheet, 46, 13, ["Rank", "Model", "Context / Price"])
    for row_idx, summary in enumerate(display_data["price_context_summary"], start=47):
        worksheet.cell(row=row_idx, column=13, value=summary["rank"])
        worksheet.cell(row=row_idx, column=14, value=summary["model"])
        worksheet.cell(row=row_idx, column=15, value=summary["value"])
    style_display_summary_range(worksheet, 47, 54, 13, 15)

    add_ranked_bar_chart(
        worksheet,
        worksheet.parent[DATA_CODING_SHEET],
        "B6",
        OPENROUTER_LEADERBOARD_BLUE,
        axis_title="Coding Arena score",
    )
    add_ranked_bar_chart(
        worksheet,
        worksheet.parent[DATA_CHAT_SHEET],
        "M6",
        OPENROUTER_ACCENT_GOLD,
        axis_title="Chat Arena score",
    )
    add_benchmark_chart(
        worksheet,
        worksheet.parent[DATA_BENCHMARKS_SHEET],
        "B32",
    )
    add_price_context_chart(
        worksheet,
        worksheet.parent[DATA_PRICE_CONTEXT_SHEET],
        "M32",
    )
    configure_display_page_setup(worksheet)


def configure_display_page_setup(worksheet) -> None:
    worksheet.print_area = "A1:Q55"
    worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A3
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_options.horizontalCentered = True
    worksheet.page_margins.left = 0.2
    worksheet.page_margins.right = 0.2
    worksheet.page_margins.top = 0.2
    worksheet.page_margins.bottom = 0.2
    worksheet.page_margins.header = 0.1
    worksheet.page_margins.footer = 0.1


def write_data_layers(
    workbook: Workbook,
    rows: list[dict[str, object]],
    generated_at: str,
    display_refreshed_at: str,
) -> list[dict[str, object]]:
    ensure_workbook_skeleton(workbook)
    display_data = prepare_display_datasets(rows)
    write_meta_sheet(
        workbook[META_SHEET],
        generated_at=generated_at,
        row_count=len(rows),
        data_refreshed_at=generated_at,
        display_refreshed_at=display_refreshed_at,
    )
    write_raw_models_sheet(workbook[RAW_MODELS_SHEET], rows)
    write_display_helper_sheets(workbook, display_data)
    return rows


def refresh_display_layer(workbook: Workbook, generated_at: str) -> None:
    ensure_workbook_skeleton(workbook)
    rows = load_rows_from_worksheet(workbook[RAW_MODELS_SHEET])
    display_data = prepare_display_datasets(rows)
    meta_values = read_meta_values(workbook)
    data_refreshed_at = meta_values.get("data_refreshed_at", generated_at)
    write_meta_sheet(
        workbook[META_SHEET],
        generated_at=generated_at,
        row_count=len(rows),
        data_refreshed_at=data_refreshed_at,
        display_refreshed_at=generated_at,
    )
    write_display_helper_sheets(workbook, display_data)
    build_display_sheet(workbook[DISPLAY_SHEET_TITLE], display_data, generated_at)


def build_workbook(
    rows: list[dict[str, object]],
    workbook_path: Path,
    generated_at: str,
    include_display: bool = True,
) -> None:
    workbook = create_workbook_skeleton()
    write_data_layers(
        workbook,
        rows,
        generated_at=generated_at,
        display_refreshed_at=generated_at if include_display else "",
    )
    if include_display:
        refresh_display_layer(workbook, generated_at)
    else:
        write_display_placeholder(workbook[DISPLAY_SHEET_TITLE], generated_at)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)


def default_paths(repo_root: Path) -> dict[str, Path]:
    output_dir = repo_root / "output" / "spreadsheet"
    cache_dir = repo_root / "tmp" / "api-cache"
    return {
        "workbook": output_dir / "llm_stats_models.xlsx",
        "all_models_csv": output_dir / "llm_stats_homepage_models.csv",
        "top20_csv": output_dir / "llm_stats_top20.csv",
        "cache_html": cache_dir / "llm_stats_homepage.html",
        "cache_json": cache_dir / "llm_stats_homepage_models.json",
    }


def fetch_and_normalize_rows(paths: dict[str, Path], generated_at: str) -> list[dict[str, object]]:
    html = load_homepage_html(paths["cache_html"])
    extracted_rows = extract_initial_homepage_models_from_html(html)
    normalized_rows = normalize_homepage_models(extracted_rows, generated_at=generated_at)
    write_json(paths["cache_json"], normalized_rows)
    write_csv(paths["all_models_csv"], normalized_rows)
    write_csv(paths["top20_csv"], normalized_rows[:20])
    return normalized_rows


def run_data_refresh(paths: dict[str, Path]) -> list[dict[str, object]]:
    generated_at = current_timestamp()
    normalized_rows = fetch_and_normalize_rows(paths, generated_at)

    if paths["workbook"].exists():
        workbook = load_workbook(paths["workbook"])
        display_refreshed_at = read_meta_values(workbook).get("display_refreshed_at", "")
        write_data_layers(
            workbook,
            normalized_rows,
            generated_at=generated_at,
            display_refreshed_at=display_refreshed_at,
        )
    else:
        workbook = create_workbook_skeleton()
        write_data_layers(
            workbook,
            normalized_rows,
            generated_at=generated_at,
            display_refreshed_at="",
        )
        write_display_placeholder(workbook[DISPLAY_SHEET_TITLE], generated_at)

    paths["workbook"].parent.mkdir(parents=True, exist_ok=True)
    workbook.save(paths["workbook"])
    return normalized_rows


def run_chart_refresh(paths: dict[str, Path]) -> None:
    if not paths["workbook"].exists():
        raise FileNotFoundError(f"Workbook not found for chart-refresh: {paths['workbook']}")
    generated_at = current_timestamp()
    workbook = load_workbook(paths["workbook"])
    refresh_display_layer(workbook, generated_at)
    workbook.save(paths["workbook"])


def run_full_refresh(paths: dict[str, Path]) -> list[dict[str, object]]:
    generated_at = current_timestamp()
    normalized_rows = fetch_and_normalize_rows(paths, generated_at)

    if paths["workbook"].exists():
        workbook = load_workbook(paths["workbook"])
        display_refreshed_at = read_meta_values(workbook).get("display_refreshed_at", "")
        write_data_layers(
            workbook,
            normalized_rows,
            generated_at=generated_at,
            display_refreshed_at=display_refreshed_at,
        )
        refresh_display_layer(workbook, generated_at)
        paths["workbook"].parent.mkdir(parents=True, exist_ok=True)
        workbook.save(paths["workbook"])
    else:
        build_workbook(
            normalized_rows,
            paths["workbook"],
            generated_at=generated_at,
            include_display=True,
        )
    return normalized_rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--mode",
        choices=["full", "data-refresh", "chart-refresh", "display-refresh"],
        default="full",
    )
    args = parser.parse_args()

    repo_root = Path(__file__).resolve().parent
    paths = default_paths(repo_root)

    if args.mode == "full":
        run_full_refresh(paths)
        return

    if args.mode == "data-refresh":
        run_data_refresh(paths)
        return

    run_chart_refresh(paths)


if __name__ == "__main__":
    main()
