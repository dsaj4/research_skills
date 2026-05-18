import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import generate_openrouter_company_tokens_excel as company_tokens


class CompanyAggregationTests(unittest.TestCase):
    def test_aggregate_company_weekly_rows_sums_model_tokens_by_provider(self) -> None:
        model_keys = ["openai/model-a", "openai/model-b", "anthropic/model-c", "Others"]
        model_rows = [
            {
                "week_start": "2026-05-04",
                "openai/model-a": 100.0,
                "openai/model-b": 50.0,
                "anthropic/model-c": 25.0,
                "Others": 10.0,
            }
        ]

        companies, rows = company_tokens.aggregate_company_weekly_rows(model_keys, model_rows)

        self.assertEqual(companies, ["anthropic", "openai", "Others"])
        self.assertEqual(rows[0]["openai"], 150.0)
        self.assertEqual(rows[0]["anthropic"], 25.0)
        self.assertEqual(rows[0]["Others"], 10.0)


class CompanyWorkbookTests(unittest.TestCase):
    def test_build_workbook_creates_workpaper_layers(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            source_csv = base / "openrouter_top_models_timeseries.csv"
            output_path = base / "openrouter_company_token_usage.xlsx"
            with source_csv.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["week_start", "openai/model-a", "openai/model-b", "anthropic/model-c", "Others"])
                writer.writerow(["2026-05-04", 100, 50, 25, 10])
                writer.writerow(["2026-05-11", 200, 75, 125, 20])

            company_tokens.build_workbook(source_csv, output_path, generated_at="2026-05-18 10:00:00 CST")

            wb = load_workbook(output_path)
            self.assertEqual(
                wb.sheetnames,
                [
                    "00_Workpaper_Index",
                    "01_Raw_Model_Weekly",
                    "02_Calc_Company_Weekly",
                    "03_Output_Latest_Rank",
                    "04_Output_Company_Share",
                    "05_Output_Weekly_Summary",
                ],
            )
            self.assertEqual(wb.active.title, "00_Workpaper_Index")
            latest = wb["03_Output_Latest_Rank"]
            self.assertEqual(latest["C7"].value, "openai")
            self.assertEqual(latest["E7"].value, 275 / 1_000_000_000)
            self.assertEqual(latest.freeze_panes, "A7")


if __name__ == "__main__":
    unittest.main()
