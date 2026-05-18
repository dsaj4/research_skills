from __future__ import annotations

import argparse
import csv
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SOURCE_URL = "https://openrouter.ai/rankings"
WORKBOOK_TITLE = "OpenRouter Model Token Usage by Company"
UNIT_RAW = "tokens"
UNIT_BILLIONS = "billions of tokens"

INDEX_SHEET = "00_Workpaper_Index"
RAW_MODEL_WEEKLY_SHEET = "01_Raw_Model_Weekly"
CALC_COMPANY_WEEKLY_SHEET = "02_Calc_Company_Weekly"
OUTPUT_LATEST_RANK_SHEET = "03_Output_Latest_Rank"
OUTPUT_COMPANY_SHARE_SHEET = "04_Output_Company_Share"
OUTPUT_WEEKLY_SUMMARY_SHEET = "05_Output_Weekly_Summary"

HEADER_ROW = 6
DATA_START_ROW = 7
MAX_OUTPUT_COMPANIES = 20


def current_timestamp() -> str:
    return datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M:%S %Z")


def format_token_amount(value: float | int | None) -> str:
    numeric = float(value or 0)
    abs_value = abs(numeric)
    if abs_value >= 1_000_000_000_000:
        return f"{numeric / 1_000_000_000_000:.2f}T tokens"
    if abs_value >= 1_000_000_000:
        return f"{numeric / 1_000_000_000:.2f}B tokens"
    if abs_value >= 1_000_000:
        return f"{numeric / 1_000_000:.2f}M tokens"
    return f"{numeric:.0f} tokens"


def load_model_weekly_rows(csv_path: Path) -> tuple[list[str], list[dict[str, float | str]]]:
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        model_keys = [header for header in headers if header != "week_start"]
        rows: list[dict[str, float | str]] = []
        for source_row in reader:
            row: dict[str, float | str] = {"week_start": source_row["week_start"]}
            for key in model_keys:
                value = source_row.get(key, "")
                row[key] = float(value) if value not in {"", None} else 0.0
            rows.append(row)
    return model_keys, rows


def company_for_model_key(model_key: str) -> str:
    if model_key == "Others":
        return "Others"
    return model_key.split("/", 1)[0] if "/" in model_key else "Unknown"


def aggregate_company_weekly_rows(model_keys: list[str], model_rows: list[dict[str, float | str]]) -> tuple[list[str], list[dict[str, float | str]]]:
    companies = sorted({company_for_model_key(key) for key in model_keys if key != "Others"})
    if "Others" in model_keys:
        companies.append("Others")

    company_rows: list[dict[str, float | str]] = []
    for model_row in model_rows:
        company_row: dict[str, float | str] = {"week_start": model_row["week_start"]}
        totals = {company: 0.0 for company in companies}
        for model_key in model_keys:
            company = company_for_model_key(model_key)
            totals[company] = totals.get(company, 0.0) + float(model_row.get(model_key, 0) or 0)
        company_row.update(totals)
        company_rows.append(company_row)

    return companies, company_rows


def latest_company_values(companies: list[str], company_rows: list[dict[str, float | str]]) -> list[tuple[str, float]]:
    if not company_rows:
        return []
    latest = company_rows[-1]
    values = [(company, float(latest.get(company, 0) or 0)) for company in companies]
    return sorted(values, key=lambda item: item[1], reverse=True)


def build_latest_rank_rows(companies: list[str], company_rows: list[dict[str, float | str]]) -> list[list[object]]:
    latest_values = latest_company_values(companies, company_rows)
    latest_week = company_rows[-1]["week_start"] if company_rows else ""
    prior_values = {}
    if len(company_rows) > 1:
        prior_values = {company: float(company_rows[-2].get(company, 0) or 0) for company in companies}

    rows: list[list[object]] = []
    for rank, (company, latest_tokens) in enumerate(latest_values[:MAX_OUTPUT_COMPANIES], start=1):
        prior_tokens = prior_values.get(company, 0)
        wow_change = (latest_tokens - prior_tokens) / prior_tokens if prior_tokens else 0
        rows.append(
            [
                rank,
                latest_week,
                company,
                latest_tokens,
                latest_tokens / 1_000_000_000,
                format_token_amount(latest_tokens),
                prior_tokens / 1_000_000_000,
                wow_change,
                SOURCE_URL,
            ]
        )
    return rows


def build_company_share_rows(companies: list[str], company_rows: list[dict[str, float | str]]) -> list[list[object]]:
    latest_values = latest_company_values(companies, company_rows)
    latest_week = company_rows[-1]["week_start"] if company_rows else ""
    total = sum(value for _, value in latest_values)
    rows: list[list[object]] = []
    for rank, (company, latest_tokens) in enumerate(latest_values[:MAX_OUTPUT_COMPANIES], start=1):
        rows.append([rank, latest_week, company, latest_tokens / 1_000_000_000, latest_tokens / total if total else 0])
    return rows


def build_weekly_summary_rows(companies: list[str], company_rows: list[dict[str, float | str]]) -> list[list[object]]:
    rows: list[list[object]] = []
    prior_total = 0.0
    for company_row in company_rows:
        week = str(company_row["week_start"])
        total = sum(float(company_row.get(company, 0) or 0) for company in companies)
        values = sorted(
            [(company, float(company_row.get(company, 0) or 0)) for company in companies],
            key=lambda item: item[1],
            reverse=True,
        )
        top_company, top_tokens = values[0] if values else ("", 0.0)
        wow_change = (total - prior_total) / prior_total if prior_total else 0
        rows.append([week, total, total / 1_000_000_000, top_company, top_tokens / 1_000_000_000, wow_change, len([c for c in companies if c != "Others"])])
        prior_total = total
    return rows


def clear_sheet(ws) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    ws._charts = []


def apply_workpaper_page_setup(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{DATA_START_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_workpaper_table(
    ws,
    sheet_title: str,
    purpose: str,
    source: str,
    generated_at: str,
    headers: list[str],
    rows: list[list[object]],
    number_formats: dict[int, str] | None = None,
    units: str = UNIT_RAW,
) -> None:
    clear_sheet(ws)
    apply_workpaper_page_setup(ws)

    dark_fill = PatternFill("solid", fgColor="0F172A")
    label_fill = PatternFill("solid", fgColor="E2E8F0")
    header_fill = PatternFill("solid", fgColor="1E293B")
    thin_gray = Side(style="thin", color="CBD5E1")
    bottom_border = Border(bottom=thin_gray)

    ws["A1"] = sheet_title
    ws["A1"].font = Font(size=15, bold=True, color="F8FAFC")
    ws["A1"].fill = dark_fill
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(6, len(headers)))

    metadata = [
        ("Purpose", purpose),
        ("Source", source),
        ("Generated at", generated_at),
        ("Unit", units),
    ]
    for row_idx, (label, value) in enumerate(metadata, start=2):
        ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row_idx, column=1).fill = label_fill
        ws.cell(row=row_idx, column=2, value=value)
        ws.cell(row=row_idx, column=2).alignment = Alignment(wrap_text=True)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header)
        cell.font = Font(bold=True, color="F8FAFC")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row_values in enumerate(rows, start=DATA_START_ROW):
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = bottom_border
            if number_formats and col_idx in number_formats:
                cell.number_format = number_formats[col_idx]
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="left")

    if rows:
        ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(headers))}{HEADER_ROW + len(rows)}"

    for col_idx, header in enumerate(headers, start=1):
        width = max(12, min(42, len(str(header)) + 4))
        if col_idx == 1:
            width = max(width, 14)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_index_sheet(
    ws,
    generated_at: str,
    source_csv_path: Path,
    model_count: int,
    company_count: int,
    week_count: int,
) -> None:
    rows = [
        ["Workbook", WORKBOOK_TITLE],
        ["Source URL", SOURCE_URL],
        ["Source CSV", str(source_csv_path)],
        ["Generated at", generated_at],
        ["Granularity", "Weekly"],
        ["Classification", "Company is derived from OpenRouter model slug prefix before '/'."],
        ["Model count", model_count],
        ["Company count", company_count],
        ["Week count", week_count],
        ["Review note", "This workbook is formatted as a workpaper: raw input, calculation layer, and output tables are separated and consistently styled."],
    ]
    write_workpaper_table(
        ws,
        "00 Workpaper Index",
        "Control sheet for source, scope, and workbook structure.",
        SOURCE_URL,
        generated_at,
        ["Item", "Value"],
        rows,
        units="Mixed",
    )
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 96


def build_workbook(source_csv_path: Path, output_path: Path, generated_at: str | None = None) -> Path:
    generated_at = generated_at or current_timestamp()
    model_keys, model_rows = load_model_weekly_rows(source_csv_path)
    companies, company_rows = aggregate_company_weekly_rows(model_keys, model_rows)

    wb = Workbook()
    wb.active.title = INDEX_SHEET
    for title in (
        RAW_MODEL_WEEKLY_SHEET,
        CALC_COMPANY_WEEKLY_SHEET,
        OUTPUT_LATEST_RANK_SHEET,
        OUTPUT_COMPANY_SHARE_SHEET,
        OUTPUT_WEEKLY_SUMMARY_SHEET,
    ):
        wb.create_sheet(title)

    write_index_sheet(wb[INDEX_SHEET], generated_at, source_csv_path, len(model_keys), len([c for c in companies if c != "Others"]), len(model_rows))

    raw_number_formats = {idx: '#,##0;[Red](#,##0);"-"' for idx in range(2, len(model_keys) + 2)}
    write_workpaper_table(
        wb[RAW_MODEL_WEEKLY_SHEET],
        "01 Raw Model Weekly",
        "Raw weekly token time series by OpenRouter model slug.",
        str(source_csv_path),
        generated_at,
        ["week_start", *model_keys],
        [[row["week_start"], *[row.get(key, 0) for key in model_keys]] for row in model_rows],
        number_formats=raw_number_formats,
    )
    wb[RAW_MODEL_WEEKLY_SHEET].column_dimensions["A"].width = 14

    company_number_formats = {idx: '#,##0;[Red](#,##0);"-"' for idx in range(2, len(companies) + 2)}
    write_workpaper_table(
        wb[CALC_COMPANY_WEEKLY_SHEET],
        "02 Calc Company Weekly",
        "Calculation layer aggregating model weekly tokens into company/provider totals.",
        RAW_MODEL_WEEKLY_SHEET,
        generated_at,
        ["week_start", *companies],
        [[row["week_start"], *[float(row.get(company, 0) or 0) for company in companies]] for row in company_rows],
        number_formats=company_number_formats,
    )
    wb[CALC_COMPANY_WEEKLY_SHEET].column_dimensions["A"].width = 14

    write_workpaper_table(
        wb[OUTPUT_LATEST_RANK_SHEET],
        "03 Output Latest Rank",
        "Latest week company ranking by OpenRouter token usage.",
        CALC_COMPANY_WEEKLY_SHEET,
        generated_at,
        ["rank", "week_start", "company", "tokens_raw", "tokens_bn", "tokens_display", "prior_week_tokens_bn", "wow_change", "source_url"],
        build_latest_rank_rows(companies, company_rows),
        number_formats={4: '#,##0;[Red](#,##0);"-"', 5: '0.00 "B"', 7: '0.00 "B"', 8: "0.00%"},
        units=UNIT_BILLIONS,
    )
    wb[OUTPUT_LATEST_RANK_SHEET].column_dimensions["C"].width = 24
    wb[OUTPUT_LATEST_RANK_SHEET].column_dimensions["I"].width = 32

    write_workpaper_table(
        wb[OUTPUT_COMPANY_SHARE_SHEET],
        "04 Output Company Share",
        "Latest week company share of tracked OpenRouter token usage.",
        CALC_COMPANY_WEEKLY_SHEET,
        generated_at,
        ["rank", "week_start", "company", "tokens_bn", "latest_share"],
        build_company_share_rows(companies, company_rows),
        number_formats={4: '0.00 "B"', 5: "0.00%"},
        units=UNIT_BILLIONS,
    )
    wb[OUTPUT_COMPANY_SHARE_SHEET].column_dimensions["C"].width = 24

    write_workpaper_table(
        wb[OUTPUT_WEEKLY_SUMMARY_SHEET],
        "05 Output Weekly Summary",
        "Weekly total token usage and top company checks.",
        CALC_COMPANY_WEEKLY_SHEET,
        generated_at,
        ["week_start", "total_tokens_raw", "total_tokens_bn", "top_company", "top_company_tokens_bn", "wow_change", "company_count"],
        build_weekly_summary_rows(companies, company_rows),
        number_formats={2: '#,##0;[Red](#,##0);"-"', 3: '0.00 "B"', 5: '0.00 "B"', 6: "0.00%"},
        units=UNIT_BILLIONS,
    )
    wb[OUTPUT_WEEKLY_SUMMARY_SHEET].column_dimensions["D"].width = 24

    wb.active = wb.sheetnames.index(INDEX_SHEET)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Generate OpenRouter company-level weekly token usage workpaper.")
    parser.add_argument("--source-csv", type=Path, default=Path("output/spreadsheet/openrouter_top_models_timeseries.csv"))
    parser.add_argument("--output-path", type=Path, default=Path("output/spreadsheet/openrouter_company_token_usage.xlsx"))
    args = parser.parse_args(argv)
    print(build_workbook(args.source_csv, args.output_path))


if __name__ == "__main__":
    main()
