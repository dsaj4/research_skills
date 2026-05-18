from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import generate_llm_stats_rankings_excel as rankings


FIXTURES_DIR = Path(__file__).parent / "fixtures"


def normalized_print_area(value: object) -> str:
    return str(value).replace("$", "").split("!")[-1].strip("'")


def sample_rows() -> list[dict[str, object]]:
    seed_rows = [
        ("claude-opus-4-6", "Claude Opus 4.6", "Anthropic", "anthropic", 14.76, 20.02, 0.913, 0.808, 0.531, 1_000_000, 5.0, 25.0, "Proprietary", "2026-02-05"),
        ("gemini-3.1-pro-preview", "Gemini 3.1 Pro", "Google", "google", 12.22, 18.62, 0.943, 0.806, 0.514, 1_048_576, 2.5, 15.0, "Proprietary", "2026-02-19"),
        ("gpt-5.4", "GPT-5.4", "OpenAI", "openai", 11.46, 16.79, 0.928, None, 0.398, 1_000_000, 2.5, 15.0, "Proprietary", "2026-03-05"),
        ("glm-5", "GLM-5", "Zhipu AI", "zai-org", 11.79, 15.95, None, 0.778, None, 200_000, 1.0, 3.2, "Open Source", "2026-02-11"),
        ("claude-opus-4-5-20251101", "Claude Opus 4.5", "Anthropic", "anthropic", 13.45, 15.80, 0.870, 0.809, 0.480, 200_000, 5.0, 25.0, "Proprietary", "2025-11-01"),
        ("seed-2.0-pro", "Seed 2.0 Pro", "ByteDance", "bytedance", 12.05, 15.40, 0.889, 0.765, None, 256_000, 1.2, 6.4, "Proprietary", "2026-02-14"),
        ("qwen3.5-27b", "Qwen3.5-27B", "Qwen", "qwen", 10.88, 14.95, 0.855, 0.724, 0.485, 262_144, 0.4, 2.2, "Open Source", "2026-02-24"),
        ("grok-4", "Grok 4", "xAI", "xai", 10.73, 14.62, 0.841, 0.701, 0.410, 256_000, 3.0, 18.0, "Proprietary", "2026-01-17"),
        ("kimi-k2.5", "Kimi K2.5", "Moonshot AI", "moonshotai", 9.88, 14.05, 0.802, 0.688, 0.365, 512_000, 1.0, 8.0, "Proprietary", "2026-01-27"),
        ("deepseek-v3.2", "DeepSeek V3.2", "DeepSeek", "deepseek", 9.55, 13.84, 0.798, 0.671, 0.352, 256_000, 0.6, 2.4, "Open Source", "2026-03-01"),
        ("claude-sonnet-4-6", "Claude Sonnet 4.6", "Anthropic", "anthropic", 9.31, 13.41, 0.846, 0.690, 0.420, 1_000_000, 3.0, 15.0, "Proprietary", "2026-02-17"),
        ("gemini-3-flash-preview", "Gemini 3 Flash Preview", "Google", "google", 9.18, 12.97, 0.781, 0.640, 0.300, 1_048_576, 0.8, 4.0, "Proprietary", "2025-12-17"),
    ]
    rows = []
    generated_at = "2026-04-01 20:00:00 CST"
    for idx, item in enumerate(seed_rows, start=1):
        (
            model_id,
            name,
            organization,
            organization_id,
            chat_score,
            coding_score,
            gpqa_score,
            swe_score,
            hle_score,
            context,
            input_price,
            output_price,
            license_label,
            announcement_date,
        ) = item
        rows.append(
            {
                "rank": idx,
                "model_id": model_id,
                "name": name,
                "organization": organization,
                "organization_id": organization_id,
                "chat_arena_score": chat_score,
                "coding_arena_score": coding_score,
                "gpqa_score": gpqa_score,
                "swe_bench_verified_score": swe_score,
                "hle_score": hle_score,
                "context": context,
                "context_display": rankings.format_context(context),
                "input_price": input_price,
                "input_price_display": rankings.format_price(input_price),
                "output_price": output_price,
                "output_price_display": rankings.format_price(output_price),
                "license": license_label,
                "announcement_date": announcement_date,
                "source_url": rankings.SOURCE_URL,
                "generated_at": generated_at,
            }
        )
    return rows


class LLMStatsExtractionTests(unittest.TestCase):
    def test_extract_initial_homepage_models_from_html(self) -> None:
        html = (FIXTURES_DIR / "homepage_fragment.html").read_text(encoding="utf-8")

        rows = rankings.extract_initial_homepage_models_from_html(html)
        normalized = rankings.normalize_homepage_models(rows, generated_at="2026-04-01 20:00:00 CST")

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["model_id"], "claude-opus-4-6")
        self.assertEqual(normalized[0]["rank"], 1)
        self.assertEqual(normalized[1]["model_id"], "gpt-5.4")
        self.assertEqual(normalized[1]["chat_arena_score"], 11.456166933572916)
        self.assertEqual(normalized[1]["license"], "Proprietary")

    def test_load_homepage_html_uses_cache_when_fetch_fails(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            cache_path = Path(tmp_dir) / "homepage.html"
            cache_path.write_text("<html>cached</html>", encoding="utf-8")

            with mock.patch.object(rankings, "fetch_homepage_html", side_effect=TimeoutError("timeout")):
                html = rankings.load_homepage_html(cache_path)

        self.assertEqual(html, "<html>cached</html>")


class WorkbookTests(unittest.TestCase):
    def test_build_workbook_creates_openrouter_style_display_layout(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "llm_stats_models.xlsx"

            rankings.build_workbook(
                sample_rows(),
                workbook_path,
                generated_at="2026-04-01 20:00:00 CST",
                include_display=True,
            )

            wb = load_workbook(workbook_path)
            self.assertEqual(
                wb.sheetnames,
                [
                    rankings.DISPLAY_SHEET_TITLE,
                    rankings.META_SHEET,
                    rankings.RAW_MODELS_SHEET,
                    rankings.DATA_TOP20_SHEET,
                    rankings.DATA_CODING_SHEET,
                    rankings.DATA_CHAT_SHEET,
                    rankings.DATA_BENCHMARKS_SHEET,
                    rankings.DATA_PRICE_CONTEXT_SHEET,
                ],
            )
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["A1"].value, rankings.DISPLAY_PAGE_TITLE)
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["B4"].value, rankings.DISPLAY_TOP_LEFT_TITLE)
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["M4"].value, rankings.DISPLAY_TOP_RIGHT_TITLE)
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["B30"].value, rankings.DISPLAY_BOTTOM_LEFT_TITLE)
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["M30"].value, rankings.DISPLAY_BOTTOM_RIGHT_TITLE)
            self.assertEqual(len(wb[rankings.DISPLAY_SHEET_TITLE]._charts), 4)
            self.assertEqual(wb[rankings.DATA_CODING_SHEET].sheet_state, "hidden")
            self.assertEqual(wb[rankings.DATA_CHAT_SHEET].sheet_state, "hidden")
            self.assertEqual(wb[rankings.RAW_MODELS_SHEET]["A2"].value, 1)
            self.assertEqual(wb[rankings.RAW_MODELS_SHEET]["C2"].value, "Claude Opus 4.6")
            self.assertEqual(
                normalized_print_area(wb[rankings.DISPLAY_SHEET_TITLE].print_area),
                "A1:Q55",
            )
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE].page_setup.orientation, wb[rankings.DISPLAY_SHEET_TITLE].ORIENTATION_LANDSCAPE)
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE].page_setup.fitToWidth, 1)
            coding_chart = wb[rankings.DISPLAY_SHEET_TITLE]._charts[0]
            chat_chart = wb[rankings.DISPLAY_SHEET_TITLE]._charts[1]
            benchmark_chart = wb[rankings.DISPLAY_SHEET_TITLE]._charts[2]
            price_context_chart = wb[rankings.DISPLAY_SHEET_TITLE]._charts[3]
            self.assertEqual(len(coding_chart.ser), 1)
            self.assertEqual(len(chat_chart.ser), 1)
            self.assertEqual(len(benchmark_chart.ser), 2)
            self.assertEqual(len(price_context_chart.ser), 1)
            self.assertIsNone(coding_chart.legend)
            self.assertGreaterEqual(coding_chart.gapWidth, 30)
            self.assertFalse(hasattr(coding_chart, "dLbls") and coding_chart.dLbls)
            self.assertEqual(wb[rankings.DATA_CODING_SHEET]["A2"].value, "Claude Opus 4.6")
            self.assertEqual(wb[rankings.DATA_CODING_SHEET]["B1"].value, "coding_arena_score")
            self.assertEqual(wb[rankings.DATA_BENCHMARKS_SHEET]["A2"].value, "GPT-5.4")
            self.assertEqual(wb[rankings.DATA_BENCHMARKS_SHEET]["B1"].value, "GPQA")
            self.assertEqual(wb[rankings.DATA_BENCHMARKS_SHEET]["C1"].value, "SWE-Bench")

    def test_data_refresh_mode_writes_placeholder_display_page(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "llm_stats_models.xlsx"

            rankings.build_workbook(
                sample_rows(),
                workbook_path,
                generated_at="2026-04-01 20:00:00 CST",
                include_display=False,
            )

            wb = load_workbook(workbook_path)
            ws = wb[rankings.DISPLAY_SHEET_TITLE]

            self.assertEqual(ws["A1"].value, rankings.DISPLAY_PAGE_TITLE)
            self.assertEqual(ws["A2"].value, rankings.DISPLAY_PLACEHOLDER_NOTE)
            self.assertIn(rankings.DISPLAY_DATA_REFRESHED_LABEL, ws["A3"].value)
            self.assertEqual(len(ws._charts), 0)
            self.assertEqual(normalized_print_area(ws.print_area), "A1:Q55")

    def test_write_csv_exports_ranked_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            csv_path = Path(tmp_dir) / "llm_stats_top20.csv"

            rankings.write_csv(csv_path, sample_rows())

            with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))

        self.assertEqual(len(rows), 12)
        self.assertEqual(rows[0]["rank"], "1")
        self.assertEqual(rows[-1]["model_id"], "gemini-3-flash-preview")
