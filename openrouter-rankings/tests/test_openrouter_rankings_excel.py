from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import generate_openrouter_rankings_excel as rankings


FIXTURES_DIR = Path(__file__).parent / "fixtures"


def sample_leaderboard_rows() -> list[dict[str, object]]:
    return [
        {
            "rank": 1,
            "model": "Model A",
            "provider": "provider-a",
            "model_url": "https://openrouter.ai/provider-a/model-a",
            "provider_url": "https://openrouter.ai/provider-a",
            "weekly_tokens_display": "120B tokens",
            "weekly_tokens_billions": 120.0,
            "wow_change_percent": 14,
            "wow_change_ratio": 0.14,
            "wow_change_direction": "up",
        },
        {
            "rank": 2,
            "model": "Model B",
            "provider": "provider-b",
            "model_url": "https://openrouter.ai/provider-b/model-b",
            "provider_url": "https://openrouter.ai/provider-b",
            "weekly_tokens_display": "110B tokens",
            "weekly_tokens_billions": 110.0,
            "wow_change_percent": -5,
            "wow_change_ratio": -0.05,
            "wow_change_direction": "down",
        },
        {
            "rank": 3,
            "model": "Kimi K2.5",
            "provider": "moonshotai",
            "model_url": "https://openrouter.ai/moonshotai/kimi-k2.5",
            "provider_url": "https://openrouter.ai/moonshotai",
            "weekly_tokens_display": "90B tokens",
            "weekly_tokens_billions": 90.0,
            "wow_change_percent": 7,
            "wow_change_ratio": 0.07,
            "wow_change_direction": "up",
        },
        {
            "rank": 4,
            "model": "Model D",
            "provider": "provider-d",
            "model_url": "https://openrouter.ai/provider-d/model-d",
            "provider_url": "https://openrouter.ai/provider-d",
            "weekly_tokens_display": "80B tokens",
            "weekly_tokens_billions": 80.0,
            "wow_change_percent": 2,
            "wow_change_ratio": 0.02,
            "wow_change_direction": "up",
        },
    ]


def sample_timeseries_payload() -> dict[str, object]:
    return {
        "forecast_key": "forecast-1w",
        "rows": [
            {
                "week_start": "2025-04-07",
                "provider-a/model-a": 80,
                "provider-b/model-b": 70,
                rankings.KIMI_SERIES_KEY: 40,
                "provider-d/model-d": 30,
                "Others": 160,
            },
            {
                "week_start": "2025-04-14",
                "provider-a/model-a": 120,
                "provider-b/model-b": 110,
                rankings.KIMI_SERIES_KEY: 90,
                "provider-d/model-d": 80,
                "Others": 220,
                "forecast-1w": 999,
            },
        ],
    }


def sample_older_timeseries_payload() -> dict[str, object]:
    return {
        "forecast_key": "forecast-legacy",
        "rows": [
            {
                "week_start": "2025-03-31",
                "provider-a/model-a": 60,
                "provider-b/model-b": 50,
                "Others": 140,
            },
            {
                "week_start": "2025-04-07",
                "provider-a/model-a": 70,
                "provider-b/model-b": 40,
                "provider-d/model-d": 20,
                "Others": 130,
            },
        ],
    }


def sample_leaderboard_history_rows() -> list[dict[str, object]]:
    return [
        {
            **sample_leaderboard_rows()[0],
            "source_url": rankings.SOURCE_URL,
            "ranking_period": rankings.RANKING_PERIOD,
            "snapshot_captured_at": "2026-03-24",
            "generated_at": "2026-03-24 09:00:00 CST",
        },
        {
            **sample_leaderboard_rows()[1],
            "source_url": rankings.SOURCE_URL,
            "ranking_period": rankings.RANKING_PERIOD,
            "snapshot_captured_at": "2026-03-24",
            "generated_at": "2026-03-24 09:00:00 CST",
        },
        {
            **sample_leaderboard_rows()[2],
            "source_url": rankings.SOURCE_URL,
            "ranking_period": rankings.RANKING_PERIOD,
            "snapshot_captured_at": "2026-03-31",
            "generated_at": "2026-03-31 09:00:00 CST",
        },
        {
            **sample_leaderboard_rows()[3],
            "source_url": rankings.SOURCE_URL,
            "ranking_period": rankings.RANKING_PERIOD,
            "snapshot_captured_at": "2026-03-31",
            "generated_at": "2026-03-31 09:00:00 CST",
        },
    ]


class TopModelsExtractionTests(unittest.TestCase):
    def test_extract_top_models_chart_payload_from_html(self) -> None:
        html = (FIXTURES_DIR / "top_models_fragment.html").read_text(encoding="utf-8")

        payload = rankings.extract_top_models_chart_payload_from_html(html)

        self.assertEqual(payload["forecast_key"], "forecast-1w")
        self.assertEqual(len(payload["rows"]), 2)
        self.assertEqual(payload["rows"][0]["week_start"], "2025-04-07")
        self.assertEqual(payload["rows"][0]["model/a"], 100)
        self.assertEqual(payload["rows"][1]["model/b"], 75)
        self.assertEqual(payload["rows"][1]["forecast-1w"], 999)

    def test_load_rankings_html_uses_cache_when_fetch_times_out(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            cache_path = Path(tmp_dir) / "rankings_page.html"
            cache_path.write_text("<html>cached</html>", encoding="utf-8")

            with mock.patch.object(rankings, "fetch_rankings_html", side_effect=TimeoutError("timeout")):
                html = rankings.load_rankings_html(cache_path)

        self.assertEqual(html, "<html>cached</html>")


class DatasetTests(unittest.TestCase):
    def test_merge_timeseries_history_keeps_older_weeks_and_replaces_overlap(self) -> None:
        merged = rankings.merge_timeseries_history(sample_older_timeseries_payload(), sample_timeseries_payload())

        self.assertEqual(merged["forecast_key"], "forecast-1w")
        self.assertEqual([row["week_start"] for row in merged["rows"]], ["2025-03-31", "2025-04-07", "2025-04-14"])
        self.assertEqual(merged["rows"][1]["provider-a/model-a"], 80)
        self.assertNotIn("provider-d/model-d", merged["rows"][0])
        self.assertEqual(merged["rows"][1]["provider-d/model-d"], 30)

    def test_load_leaderboard_rows_from_workbook_can_filter_to_latest_snapshot(self) -> None:
        wb = rankings.create_workbook_skeleton()
        rankings.write_raw_leaderboard_sheet(wb[rankings.RAW_LEADERBOARD_SHEET], sample_leaderboard_history_rows())

        latest_rows = rankings.load_leaderboard_rows_from_workbook(wb, latest_snapshot_only=True)

        self.assertEqual(len(latest_rows), 2)
        self.assertEqual({row["snapshot_captured_at"] for row in latest_rows}, {"2026-03-31"})
        self.assertEqual([row["rank"] for row in latest_rows], [3, 4])

    def test_kimi_emphasis_series_color_map_highlights_kimi_and_mutes_others(self) -> None:
        series_keys = [
            "provider-a/model-a",
            rankings.KIMI_SERIES_KEY,
            "Others",
            "provider-b/model-b",
            "provider-c/model-c",
        ]

        color_map = rankings.kimi_emphasis_series_color_map(series_keys)

        self.assertEqual(color_map[rankings.KIMI_SERIES_KEY], rankings.KIMI_HIGHLIGHT_COLOR)
        self.assertEqual(color_map["Others"], rankings.KIMI_MUTED_COLOR)
        self.assertEqual(color_map["provider-a/model-a"], rankings.KIMI_REFERENCE_PALETTE[1])
        self.assertNotEqual(color_map["provider-a/model-a"], rankings.KIMI_HIGHLIGHT_COLOR)

    def test_prepare_display_datasets_keeps_top_models_kimi_and_others(self) -> None:
        display_data = rankings.prepare_display_datasets(sample_timeseries_payload(), sample_leaderboard_rows())

        self.assertEqual(display_data["trend_headers"], ["week_start", "Model A", "Model B", "Kimi K2.5", "Model D", "Others"])
        self.assertEqual(display_data["legend_entries"][-1]["label"], "Kimi K2.5")
        self.assertEqual(display_data["legend_entries"][-1]["color"], rankings.KIMI_HIGHLIGHT_COLOR)
        self.assertEqual(display_data["leaderboard_rows"][0]["model"], "Model A")
        self.assertEqual(display_data["kimi_summary"][0]["label"], "\u6700\u65b0\u5468")


class WorkbookStructureTests(unittest.TestCase):
    def test_build_workbook_creates_data_layers_and_display_page(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "openrouter_top_models.xlsx"

            rankings.build_workbook(
                sample_timeseries_payload(),
                sample_leaderboard_rows(),
                workbook_path,
                generated_at="2026-04-01 12:00:00 CST",
            )

            wb = load_workbook(workbook_path)
            self.assertEqual(
                wb.sheetnames,
                [rankings.DISPLAY_SHEET_TITLE, "Meta", "Raw_TimeSeries", "Raw_Leaderboard", "Data_Trend", "Data_Leaderboard", "Data_Kimi", "Data_WoW"],
            )
            self.assertNotIn("Charts", wb.sheetnames)
            self.assertEqual(wb["Data_Trend"].sheet_state, "hidden")
            self.assertEqual(wb["Data_Leaderboard"].sheet_state, "hidden")
            self.assertEqual(wb["Data_Kimi"].sheet_state, "hidden")
            self.assertEqual(wb["Data_WoW"].sheet_state, "hidden")
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["A1"].value, rankings.DISPLAY_PAGE_TITLE)
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["B4"].value, rankings.DISPLAY_TOP_LEFT_TITLE)
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["M4"].value, rankings.DISPLAY_TOP_RIGHT_TITLE)
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["B30"].value, rankings.DISPLAY_BOTTOM_LEFT_TITLE)
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE]["M30"].value, rankings.DISPLAY_BOTTOM_RIGHT_TITLE)
            self.assertEqual(len(wb[rankings.DISPLAY_SHEET_TITLE]._charts), 4)
            top_chart = wb[rankings.DISPLAY_SHEET_TITLE]._charts[0]
            self.assertEqual(top_chart.__class__.__name__, "BarChart")
            self.assertIsNotNone(top_chart.legend)
            self.assertEqual(top_chart.legend.position, "r")
            self.assertEqual(top_chart.gapWidth, 0)
            self.assertEqual(top_chart.overlap, 100.0)
            self.assertEqual(top_chart.ser[0].cat.numRef.f, "'Data_Trend'!$A$2:$A$3")
            self.assertEqual(top_chart.ser[0].val.numRef.f, "'Data_Trend'!$B$2:$B$3")
            self.assertEqual(wb[rankings.DISPLAY_SHEET_TITLE].print_area, "'展示页'!$A$1:$Q$55")
            raw_headers = [cell.value for cell in wb["Raw_TimeSeries"][1]]
            self.assertIn("forecast-1w", raw_headers)


class ModeTests(unittest.TestCase):
    def test_data_refresh_appends_incremental_history(self) -> None:
        fixture_html = (FIXTURES_DIR / "top_models_fragment.html").read_text(encoding="utf-8")

        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            workbook_path = base / "openrouter_top_models.xlsx"
            leaderboard_csv = base / "openrouter_top_models.csv"
            timeseries_csv = base / "openrouter_top_models_timeseries.csv"
            cache_path = base / "rankings_page.html"
            cache_path.write_text(fixture_html, encoding="utf-8")

            rankings.build_workbook(
                sample_older_timeseries_payload(),
                sample_leaderboard_history_rows(),
                workbook_path,
                generated_at="2026-03-31 09:00:00 CST",
            )

            with mock.patch.object(rankings, "fetch_rankings_html", side_effect=TimeoutError("timeout")):
                with mock.patch.object(rankings, "current_timestamp", return_value="2026-04-07 09:00:00 CST"):
                    with mock.patch.object(rankings, "LEADERBOARD_SNAPSHOT", sample_leaderboard_rows()):
                        rankings.main(
                            [
                                "--mode",
                                "data-refresh",
                                "--workbook-path",
                                str(workbook_path),
                                "--leaderboard-csv-path",
                                str(leaderboard_csv),
                                "--timeseries-csv-path",
                                str(timeseries_csv),
                                "--cache-path",
                                str(cache_path),
                            ]
                        )

            refreshed = load_workbook(workbook_path)
            merged_timeseries = rankings.load_timeseries_payload_from_workbook(refreshed)
            self.assertEqual([row["week_start"] for row in merged_timeseries["rows"]], ["2025-03-31", "2025-04-07", "2025-04-14"])

            leaderboard_history = rankings.load_leaderboard_rows_from_workbook(refreshed, latest_snapshot_only=False)
            self.assertEqual({row["snapshot_captured_at"] for row in leaderboard_history}, {"2026-03-24", "2026-03-31", "2026-04-07"})
            latest_snapshot_rows = rankings.load_leaderboard_rows_from_workbook(refreshed, latest_snapshot_only=True)
            self.assertEqual(len(latest_snapshot_rows), 4)
            self.assertEqual({row["snapshot_captured_at"] for row in latest_snapshot_rows}, {"2026-04-07"})

    def test_full_mode_appends_incremental_history(self) -> None:
        fixture_html = (FIXTURES_DIR / "top_models_fragment.html").read_text(encoding="utf-8")

        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            workbook_path = base / "openrouter_top_models.xlsx"
            leaderboard_csv = base / "openrouter_top_models.csv"
            timeseries_csv = base / "openrouter_top_models_timeseries.csv"
            cache_path = base / "rankings_page.html"
            cache_path.write_text(fixture_html, encoding="utf-8")

            rankings.build_workbook(
                sample_older_timeseries_payload(),
                sample_leaderboard_history_rows(),
                workbook_path,
                generated_at="2026-03-31 09:00:00 CST",
            )

            with mock.patch.object(rankings, "fetch_rankings_html", side_effect=TimeoutError("timeout")):
                with mock.patch.object(rankings, "current_timestamp", return_value="2026-04-07 09:00:00 CST"):
                    with mock.patch.object(rankings, "LEADERBOARD_SNAPSHOT", sample_leaderboard_rows()):
                        rankings.main(
                            [
                                "--mode",
                                "full",
                                "--workbook-path",
                                str(workbook_path),
                                "--leaderboard-csv-path",
                                str(leaderboard_csv),
                                "--timeseries-csv-path",
                                str(timeseries_csv),
                                "--cache-path",
                                str(cache_path),
                            ]
                        )

            refreshed = load_workbook(workbook_path)
            merged_timeseries = rankings.load_timeseries_payload_from_workbook(refreshed)
            self.assertEqual([row["week_start"] for row in merged_timeseries["rows"]], ["2025-03-31", "2025-04-07", "2025-04-14"])
            latest_snapshot_rows = rankings.load_leaderboard_rows_from_workbook(refreshed, latest_snapshot_only=True)
            self.assertEqual({row["snapshot_captured_at"] for row in latest_snapshot_rows}, {"2026-04-07"})

    def test_data_refresh_preserves_existing_display_page(self) -> None:
        fixture_html = (FIXTURES_DIR / "top_models_fragment.html").read_text(encoding="utf-8")

        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            workbook_path = base / "openrouter_top_models.xlsx"
            leaderboard_csv = base / "openrouter_top_models.csv"
            timeseries_csv = base / "openrouter_top_models_timeseries.csv"
            cache_path = base / "rankings_page.html"
            cache_path.write_text(fixture_html, encoding="utf-8")

            with mock.patch.object(rankings, "fetch_rankings_html", side_effect=TimeoutError("timeout")):
                rankings.main(
                    [
                        "--mode",
                        "full",
                        "--workbook-path",
                        str(workbook_path),
                        "--leaderboard-csv-path",
                        str(leaderboard_csv),
                        "--timeseries-csv-path",
                        str(timeseries_csv),
                        "--cache-path",
                        str(cache_path),
                    ]
                )

            wb = load_workbook(workbook_path)
            wb[rankings.DISPLAY_SHEET_TITLE]["A1"] = "display sentinel"
            wb.save(workbook_path)

            with mock.patch.object(rankings, "fetch_rankings_html", side_effect=TimeoutError("timeout")):
                rankings.main(
                    [
                        "--mode",
                        "data-refresh",
                        "--workbook-path",
                        str(workbook_path),
                        "--leaderboard-csv-path",
                        str(leaderboard_csv),
                        "--timeseries-csv-path",
                        str(timeseries_csv),
                        "--cache-path",
                        str(cache_path),
                    ]
                )

            refreshed = load_workbook(workbook_path)
            self.assertEqual(refreshed[rankings.DISPLAY_SHEET_TITLE]["A1"].value, "display sentinel")
            self.assertEqual(len(refreshed[rankings.DISPLAY_SHEET_TITLE]._charts), 4)
            self.assertFalse((base / "openrouter_top_models_kimi_emphasis.xlsx").exists())

    def test_display_refresh_rebuilds_display_without_fetching(self) -> None:
        fixture_html = (FIXTURES_DIR / "top_models_fragment.html").read_text(encoding="utf-8")

        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            workbook_path = base / "openrouter_top_models.xlsx"
            leaderboard_csv = base / "openrouter_top_models.csv"
            timeseries_csv = base / "openrouter_top_models_timeseries.csv"
            cache_path = base / "rankings_page.html"
            cache_path.write_text(fixture_html, encoding="utf-8")

            with mock.patch.object(rankings, "fetch_rankings_html", side_effect=TimeoutError("timeout")):
                rankings.main(
                    [
                        "--mode",
                        "full",
                        "--workbook-path",
                        str(workbook_path),
                        "--leaderboard-csv-path",
                        str(leaderboard_csv),
                        "--timeseries-csv-path",
                        str(timeseries_csv),
                        "--cache-path",
                        str(cache_path),
                    ]
                )

            wb = load_workbook(workbook_path)
            ws = wb[rankings.DISPLAY_SHEET_TITLE]
            ws["A1"] = "broken display"
            ws._charts = []
            wb.save(workbook_path)

            with mock.patch.object(rankings, "fetch_rankings_html", side_effect=AssertionError("display-refresh should not fetch")):
                rankings.main(
                    [
                        "--mode",
                        "display-refresh",
                        "--workbook-path",
                        str(workbook_path),
                        "--leaderboard-csv-path",
                        str(leaderboard_csv),
                        "--timeseries-csv-path",
                        str(timeseries_csv),
                        "--cache-path",
                        str(cache_path),
                    ]
                )

            rebuilt = load_workbook(workbook_path)
            self.assertEqual(rebuilt[rankings.DISPLAY_SHEET_TITLE]["A1"].value, rankings.DISPLAY_PAGE_TITLE)
            self.assertEqual(len(rebuilt[rankings.DISPLAY_SHEET_TITLE]._charts), 4)

    def test_smoke_modes_write_expected_outputs(self) -> None:
        fixture_html = (FIXTURES_DIR / "top_models_fragment.html").read_text(encoding="utf-8")

        with tempfile.TemporaryDirectory() as tmp_dir:
            base = Path(tmp_dir)
            workbook_path = base / "openrouter_top_models.xlsx"
            leaderboard_csv = base / "openrouter_top_models.csv"
            timeseries_csv = base / "openrouter_top_models_timeseries.csv"
            cache_path = base / "rankings_page.html"
            cache_path.write_text(fixture_html, encoding="utf-8")

            for mode in ("full", "data-refresh", "display-refresh"):
                if mode == "display-refresh" and not workbook_path.exists():
                    with mock.patch.object(rankings, "fetch_rankings_html", side_effect=TimeoutError("timeout")):
                        rankings.main(
                            [
                                "--mode",
                                "data-refresh",
                                "--workbook-path",
                                str(workbook_path),
                                "--leaderboard-csv-path",
                                str(leaderboard_csv),
                                "--timeseries-csv-path",
                                str(timeseries_csv),
                                "--cache-path",
                                str(cache_path),
                            ]
                        )

                with mock.patch.object(rankings, "fetch_rankings_html", side_effect=TimeoutError("timeout")):
                    rankings.main(
                        [
                            "--mode",
                            mode,
                            "--workbook-path",
                            str(workbook_path),
                            "--leaderboard-csv-path",
                            str(leaderboard_csv),
                            "--timeseries-csv-path",
                            str(timeseries_csv),
                            "--cache-path",
                            str(cache_path),
                        ]
                    )

            self.assertTrue(workbook_path.exists())
            self.assertTrue(leaderboard_csv.exists())
            self.assertTrue(timeseries_csv.exists())


class SkillStructureTests(unittest.TestCase):
    def test_repo_local_skills_use_standard_structure(self) -> None:
        root = Path(rankings.__file__).parent
        expected_files = [
            root / "skills" / "openrouter-rankings-data-refresh" / "SKILL.md",
            root / "skills" / "openrouter-rankings-data-refresh" / "agents" / "openai.yaml",
            root / "skills" / "openrouter-rankings-data-refresh" / "references" / "workflow.md",
            root / "skills" / "openrouter-rankings-data-refresh" / "scripts" / "run_data_refresh.py",
            root / "skills" / "openrouter-rankings-data-refresh" / "evals" / "evals.json",
            root / "skills" / "openrouter-rankings-display-refresh" / "SKILL.md",
            root / "skills" / "openrouter-rankings-display-refresh" / "agents" / "openai.yaml",
            root / "skills" / "openrouter-rankings-display-refresh" / "references" / "layout-reference.md",
            root / "skills" / "openrouter-rankings-display-refresh" / "scripts" / "run_display_refresh.py",
            root / "skills" / "openrouter-rankings-display-refresh" / "evals" / "evals.json",
        ]

        for path in expected_files:
            self.assertTrue(path.exists(), msg=str(path))


if __name__ == "__main__":
    unittest.main()
