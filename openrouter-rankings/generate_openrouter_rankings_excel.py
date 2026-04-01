from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import datetime
from pathlib import Path
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE_URL = "https://openrouter.ai/rankings"
RANKING_PERIOD = "This Week"
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/146.0.0.0 Safari/537.36"
LEADERBOARD_SNAPSHOT_CAPTURED_AT = "2026-03-31"
DISPLAY_SHEET_TITLE = "\u5c55\u793a\u9875"
DATA_TREND_SHEET = "Data_Trend"
DATA_LEADERBOARD_SHEET = "Data_Leaderboard"
DATA_KIMI_SHEET = "Data_Kimi"
DATA_WOW_SHEET = "Data_WoW"
RAW_TIMESERIES_SHEET = "Raw_TimeSeries"
RAW_LEADERBOARD_SHEET = "Raw_Leaderboard"
META_SHEET = "Meta"
DISPLAY_TREND_MAX_MODELS = 4
MAX_TREND_ROWS = 200
MAX_LEADERBOARD_ROWS = 20
MAX_KIMI_ROWS = 200
MAX_WOW_ROWS = 20
DISPLAY_PAGE_TITLE = "OpenRouter \u56fe\u8868\u603b\u89c8"
DISPLAY_PAGE_SUBTITLE = "\u56fe\u8868\u6765\u81ea\u5185\u90e8\u6570\u636e\u5c42\uff1b\u5e03\u5c40\u53c2\u8003 example-charts.xlsx\uff0c\u5c55\u793a\u5c42\u53ea\u627f\u8f7d\u6458\u8981\u4e0e\u56fe\u8868\u3002"
DISPLAY_TOP_LEFT_TITLE = "[\u70ed\u95e8\u6a21\u578b-\u5468\u8d8b\u52bf] Top Models Weekly Trend"
DISPLAY_TOP_RIGHT_TITLE = "[\u6392\u884c\u699c-\u6700\u65b0] Latest Leaderboard"
DISPLAY_BOTTOM_LEFT_TITLE = "[Kimi-\u805a\u7126] Kimi Spotlight"
DISPLAY_BOTTOM_RIGHT_TITLE = "[\u6392\u884c\u699c-WoW] WoW Change"
DISPLAY_REFRESHED_AT_LABEL = "\u5c55\u793a\u5237\u65b0\u65f6\u95f4\uff1a"
DISPLAY_DATA_REFRESHED_LABEL = "\u6700\u8fd1\u6570\u636e\u5237\u65b0\uff1a"
DISPLAY_PLACEHOLDER_NOTE = "\u6570\u636e\u5c42\u5df2\u5237\u65b0\u3002\u8fd0\u884c display-refresh \u4ee5\u91cd\u5efa\u5c55\u793a\u9875\u3002"
DISPLAY_LATEST_WEEK_LABEL = "\u6700\u8fd1\u4e00\u5468\uff1a"
DISPLAY_SNAPSHOT_DATE_LABEL = "\u5feb\u7167\u65e5\u671f\uff1a"
DISPLAY_KIMI_METRICS_TITLE = "Kimi \u6307\u6807"
DISPLAY_WOW_SUMMARY_TITLE = "\u6392\u884c\u699c WoW \u6458\u8981"
SERIES_FALLBACK_COLORS = [
    "0088FE",
    "00C49F",
    "FFBB28",
    "FF8042",
    "FF6347",
    "4682B4",
    "9ACD32",
    "DA70D6",
    "40E0D0",
    "FF69B4",
    "DAA520",
    "7B68EE",
    "F08080",
    "6B8E23",
    "DB7093",
    "3CB371",
]
KIMI_SERIES_KEY = "moonshotai/kimi-k2.5-0127"
KIMI_REFERENCE_PALETTE = [
    "196BA5",
    "E4C864",
    "B6B6B6",
    "7699C8",
    "B9CDE5",
    "948A54",
    "7F7F7F",
    "FAC090",
    "F1EADA",
    "98CEDD",
]
KIMI_HIGHLIGHT_COLOR = KIMI_REFERENCE_PALETTE[0]
KIMI_MUTED_COLOR = KIMI_REFERENCE_PALETTE[8]
KIMI_PALETTE_SOURCE_NOTE = "User-provided 10-color RGB palette reference from 2026-03-31."
KIMI_LEADERBOARD_COLOR = KIMI_REFERENCE_PALETTE[3]
KIMI_WOW_COLOR = KIMI_REFERENCE_PALETTE[5]

LEADERBOARD_SNAPSHOT = [
    {"rank": 1, "model": "MiMo-V2-Pro", "provider": "xiaomi", "model_url": "https://openrouter.ai/xiaomi/mimo-v2-pro", "provider_url": "https://openrouter.ai/xiaomi", "weekly_tokens_display": "4.19T tokens", "wow_change_percent": 112, "wow_change_direction": "up"},
    {"rank": 2, "model": "Step 3.5 Flash (free)", "provider": "stepfun", "model_url": "https://openrouter.ai/stepfun/step-3.5-flash:free", "provider_url": "https://openrouter.ai/stepfun", "weekly_tokens_display": "1.47T tokens", "wow_change_percent": -4, "wow_change_direction": "down"},
    {"rank": 3, "model": "MiniMax M2.7", "provider": "minimax", "model_url": "https://openrouter.ai/minimax/minimax-m2.7", "provider_url": "https://openrouter.ai/minimax", "weekly_tokens_display": "1.31T tokens", "wow_change_percent": 148, "wow_change_direction": "up"},
    {"rank": 4, "model": "DeepSeek V3.2", "provider": "deepseek", "model_url": "https://openrouter.ai/deepseek/deepseek-v3.2", "provider_url": "https://openrouter.ai/deepseek", "weekly_tokens_display": "1.24T tokens", "wow_change_percent": 8, "wow_change_direction": "up"},
    {"rank": 5, "model": "Claude Sonnet 4.6", "provider": "anthropic", "model_url": "https://openrouter.ai/anthropic/claude-sonnet-4.6", "provider_url": "https://openrouter.ai/anthropic", "weekly_tokens_display": "1.04T tokens", "wow_change_percent": 1, "wow_change_direction": "up"},
    {"rank": 6, "model": "Claude Opus 4.6", "provider": "anthropic", "model_url": "https://openrouter.ai/anthropic/claude-opus-4.6", "provider_url": "https://openrouter.ai/anthropic", "weekly_tokens_display": "987B tokens", "wow_change_percent": 0, "wow_change_direction": "down"},
    {"rank": 7, "model": "Gemini 3 Flash Preview", "provider": "google", "model_url": "https://openrouter.ai/google/gemini-3-flash-preview", "provider_url": "https://openrouter.ai/google", "weekly_tokens_display": "976B tokens", "wow_change_percent": 5, "wow_change_direction": "up"},
    {"rank": 8, "model": "GLM 5 Turbo", "provider": "z-ai", "model_url": "https://openrouter.ai/z-ai/glm-5-turbo", "provider_url": "https://openrouter.ai/z-ai", "weekly_tokens_display": "966B tokens", "wow_change_percent": -13, "wow_change_direction": "down"},
    {"rank": 9, "model": "MiniMax M2.5", "provider": "minimax", "model_url": "https://openrouter.ai/minimax/minimax-m2.5", "provider_url": "https://openrouter.ai/minimax", "weekly_tokens_display": "836B tokens", "wow_change_percent": -32, "wow_change_direction": "down"},
    {"rank": 10, "model": "Grok 4.1 Fast", "provider": "x-ai", "model_url": "https://openrouter.ai/x-ai/grok-4.1-fast", "provider_url": "https://openrouter.ai/x-ai", "weekly_tokens_display": "628B tokens", "wow_change_percent": 33, "wow_change_direction": "up"},
    {"rank": 11, "model": "Gemini 2.5 Flash Lite", "provider": "google", "model_url": "https://openrouter.ai/google/gemini-2.5-flash-lite", "provider_url": "https://openrouter.ai/google", "weekly_tokens_display": "569B tokens", "wow_change_percent": 14, "wow_change_direction": "up"},
    {"rank": 12, "model": "MiMo-V2-Omni", "provider": "xiaomi", "model_url": "https://openrouter.ai/xiaomi/mimo-v2-omni", "provider_url": "https://openrouter.ai/xiaomi", "weekly_tokens_display": "565B tokens", "wow_change_percent": 101, "wow_change_direction": "up"},
    {"rank": 13, "model": "Nemotron 3 Super (free)", "provider": "nvidia", "model_url": "https://openrouter.ai/nvidia/nemotron-3-super-120b-a12b:free", "provider_url": "https://openrouter.ai/nvidia", "weekly_tokens_display": "565B tokens", "wow_change_percent": 9, "wow_change_direction": "up"},
    {"rank": 14, "model": "Gemini 2.5 Flash", "provider": "google", "model_url": "https://openrouter.ai/google/gemini-2.5-flash", "provider_url": "https://openrouter.ai/google", "weekly_tokens_display": "534B tokens", "wow_change_percent": -8, "wow_change_direction": "down"},
    {"rank": 15, "model": "Kimi K2.5", "provider": "moonshotai", "model_url": "https://openrouter.ai/moonshotai/kimi-k2.5", "provider_url": "https://openrouter.ai/moonshotai", "weekly_tokens_display": "506B tokens", "wow_change_percent": -9, "wow_change_direction": "down"},
    {"rank": 16, "model": "gpt-oss-120b", "provider": "openai", "model_url": "https://openrouter.ai/openai/gpt-oss-120b", "provider_url": "https://openrouter.ai/openai", "weekly_tokens_display": "466B tokens", "wow_change_percent": 12, "wow_change_direction": "up"},
    {"rank": 17, "model": "GLM 5", "provider": "z-ai", "model_url": "https://openrouter.ai/z-ai/glm-5", "provider_url": "https://openrouter.ai/z-ai", "weekly_tokens_display": "380B tokens", "wow_change_percent": 51, "wow_change_direction": "up"},
    {"rank": 18, "model": "GPT-5.4", "provider": "openai", "model_url": "https://openrouter.ai/openai/gpt-5.4", "provider_url": "https://openrouter.ai/openai", "weekly_tokens_display": "339B tokens", "wow_change_percent": 26, "wow_change_direction": "up"},
    {"rank": 19, "model": "Claude Sonnet 4.5", "provider": "anthropic", "model_url": "https://openrouter.ai/anthropic/claude-sonnet-4.5", "provider_url": "https://openrouter.ai/anthropic", "weekly_tokens_display": "338B tokens", "wow_change_percent": -2, "wow_change_direction": "down"},
    {"rank": 20, "model": "Claude Haiku 4.5", "provider": "anthropic", "model_url": "https://openrouter.ai/anthropic/claude-haiku-4.5", "provider_url": "https://openrouter.ai/anthropic", "weekly_tokens_display": "281B tokens", "wow_change_percent": -1, "wow_change_direction": "down"},
]

TOP_CHART_SERIES_SNAPSHOT = [
    {"dataKey": "anthropic/claude-3-7-sonnet-20250219", "fill": "#0088FE"},
    {"dataKey": "google/gemini-2.0-flash-001", "fill": "#00C49F"},
    {"dataKey": "openai/gpt-4o-mini", "fill": "#FFBB28"},
    {"dataKey": "openrouter/quasar-alpha", "fill": "#FF8042"},
    {"dataKey": "google/gemini-2.5-pro-exp-03-25:free", "fill": "#FF6347"},
    {"dataKey": "openrouter/optimus-alpha", "fill": "#4682B4"},
    {"dataKey": "deepseek/deepseek-chat-v3-0324:free", "fill": "#9ACD32"},
    {"dataKey": "google/gemini-2.5-pro-preview-03-25", "fill": "#DA70D6"},
    {"dataKey": "anthropic/claude-3-7-sonnet-20250219:thinking", "fill": "#40E0D0"},
    {"dataKey": "Others", "fill": "#FF69B4"},
    {"dataKey": "deepseek/deepseek-chat-v3-0324", "fill": "#DAA520"},
    {"dataKey": "openai/gpt-4.1-2025-04-14", "fill": "#7B68EE"},
    {"dataKey": "google/gemini-flash-1.5-8b", "fill": "#F08080"},
    {"dataKey": "google/gemini-2.5-flash-preview-04-17", "fill": "#6B8E23"},
    {"dataKey": "deepseek/deepseek-r1:free", "fill": "#DB7093"},
    {"dataKey": "google/gemini-2.5-pro-exp-03-25", "fill": "#3CB371"},
    {"dataKey": "meta-llama/llama-3.3-70b-instruct", "fill": "#BDB76B"},
    {"dataKey": "google/gemini-2.5-flash-preview-04-17:thinking", "fill": "#800080"},
    {"dataKey": "anthropic/claude-4-sonnet-20250522", "fill": "#FF4500"},
    {"dataKey": "google/gemini-2.5-flash-preview-05-20", "fill": "#2E8B57"},
    {"dataKey": "google/gemini-2.5-pro-preview-06-05", "fill": "#0088FE"},
    {"dataKey": "google/gemini-2.0-flash-lite-001", "fill": "#00C49F"},
    {"dataKey": "google/gemini-2.5-flash-lite-preview-06-17", "fill": "#FFBB28"},
    {"dataKey": "google/gemini-2.5-flash", "fill": "#FF8042"},
    {"dataKey": "google/gemini-2.5-pro", "fill": "#FF6347"},
    {"dataKey": "deepseek/deepseek-r1-0528:free", "fill": "#4682B4"},
    {"dataKey": "qwen/qwen3-coder-480b-a35b-07-25:free", "fill": "#9ACD32"},
    {"dataKey": "qwen/qwen3-coder-480b-a35b-07-25", "fill": "#DA70D6"},
    {"dataKey": "openrouter/horizon-beta", "fill": "#40E0D0"},
    {"dataKey": "x-ai/grok-code-fast-1", "fill": "#FF69B4"},
    {"dataKey": "deepseek/deepseek-chat-v3.1", "fill": "#DAA520"},
    {"dataKey": "qwen/qwen3-30b-a3b-04-28", "fill": "#7B68EE"},
    {"dataKey": "openai/gpt-4.1-mini-2025-04-14", "fill": "#F08080"},
    {"dataKey": "deepseek/deepseek-chat-v3.1:free", "fill": "#6B8E23"},
    {"dataKey": "openrouter/sonoma-sky-alpha", "fill": "#DB7093"},
    {"dataKey": "x-ai/grok-4-fast:free", "fill": "#3CB371"},
    {"dataKey": "openai/gpt-5-2025-08-07", "fill": "#BDB76B"},
    {"dataKey": "anthropic/claude-4.5-sonnet-20250929", "fill": "#800080"},
    {"dataKey": "google/gemini-2.5-flash-lite", "fill": "#FF4500"},
    {"dataKey": "x-ai/grok-4-fast", "fill": "#2E8B57"},
    {"dataKey": "openai/gpt-oss-20b", "fill": "#0088FE"},
    {"dataKey": "qwen/qwen3-coder-30b-a3b-instruct", "fill": "#00C49F"},
    {"dataKey": "tngtech/deepseek-r1t2-chimera:free", "fill": "#FFBB28"},
    {"dataKey": "minimax/minimax-m2:free", "fill": "#FF8042"},
    {"dataKey": "minimax/minimax-m2", "fill": "#FF6347"},
    {"dataKey": "openrouter/polaris-alpha", "fill": "#4682B4"},
    {"dataKey": "x-ai/grok-4.1-fast", "fill": "#9ACD32"},
    {"dataKey": "openrouter/sherlock-think-alpha", "fill": "#DA70D6"},
    {"dataKey": "x-ai/grok-4.1-fast:free", "fill": "#40E0D0"},
    {"dataKey": "google/gemini-3-pro-preview-20251117", "fill": "#FF69B4"},
    {"dataKey": "anthropic/claude-4.5-opus-20251124", "fill": "#DAA520"},
    {"dataKey": "openai/gpt-oss-120b", "fill": "#7B68EE"},
    {"dataKey": "deepseek/deepseek-v3.2-20251201", "fill": "#F08080"},
    {"dataKey": "xiaomi/mimo-v2-flash-20251210:free", "fill": "#6B8E23"},
    {"dataKey": "google/gemini-3-flash-preview-20251217", "fill": "#DB7093"},
    {"dataKey": "moonshotai/kimi-k2.5-0127", "fill": "#3CB371"},
    {"dataKey": "minimax/minimax-m2.1", "fill": "#BDB76B"},
    {"dataKey": "minimax/minimax-m2.5-20260211", "fill": "#800080"},
    {"dataKey": "z-ai/glm-5-20260211", "fill": "#FF4500"},
    {"dataKey": "anthropic/claude-4.6-opus-20260205", "fill": "#2E8B57"},
    {"dataKey": "arcee-ai/trinity-large-preview:free", "fill": "#0088FE"},
    {"dataKey": "anthropic/claude-4.6-sonnet-20260217", "fill": "#00C49F"},
    {"dataKey": "stepfun/step-3.5-flash:free", "fill": "#FFBB28"},
    {"dataKey": "openrouter/hunter-alpha", "fill": "#FF8042"},
    {"dataKey": "xiaomi/mimo-v2-pro-20260318", "fill": "#FF6347"},
    {"dataKey": "z-ai/glm-5-turbo-20260315", "fill": "#4682B4"},
    {"dataKey": "minimax/minimax-m2.7-20260318", "fill": "#9ACD32"},
]


def current_timestamp() -> str:
    return datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M:%S %Z")


def snapshot_date_from_generated_at(generated_at: str) -> str:
    match = re.match(r"(\d{4}-\d{2}-\d{2})", generated_at)
    if match:
        return match.group(1)
    return LEADERBOARD_SNAPSHOT_CAPTURED_AT


def tokens_to_billions(display_value: str) -> float:
    value = display_value.replace(" tokens", "").strip().upper()
    if value.endswith("T"):
        return round(float(value[:-1]) * 1000, 2)
    if value.endswith("B"):
        return round(float(value[:-1]), 2)
    raise ValueError(f"Unsupported token unit: {display_value}")


def format_token_amount(value: float | int | None) -> str:
    numeric = float(value or 0)
    abs_value = abs(numeric)
    if abs_value >= 1_000_000_000_000:
        return f"{numeric / 1_000_000_000_000:.2f}T tokens"
    if abs_value >= 1_000_000_000:
        return f"{numeric / 1_000_000_000:.2f}B tokens"
    if abs_value >= 1_000_000:
        return f"{numeric / 1_000_000:.2f}M tokens"
    return f"{numeric:.0f} tokens"


def format_percent(value: float | int | None) -> str:
    return f"{float(value or 0):.2f}%"


def fetch_rankings_html(url: str = SOURCE_URL) -> str:
    request = Request(url, headers={"User-Agent": USER_AGENT})
    with urlopen(request, timeout=60) as response:
        return response.read().decode("utf-8")


def load_rankings_html(cache_path: Path) -> str:
    try:
        html = fetch_rankings_html()
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(html, encoding="utf-8")
        return html
    except Exception:
        if cache_path.exists():
            return cache_path.read_text(encoding="utf-8")
        raise


def extract_balanced_block(text: str, start_idx: int, open_char: str = "[", close_char: str = "]") -> str:
    depth = 0
    in_string = False
    escaped = False
    for idx in range(start_idx, len(text)):
        char = text[idx]
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if char == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if char == open_char:
            depth += 1
        elif char == close_char:
            depth -= 1
            if depth == 0:
                return text[start_idx : idx + 1]
    raise ValueError("Unbalanced block in rankings HTML")


def extract_top_models_chart_payload_from_html(html: str) -> dict[str, object]:
    forecast_match = re.search(r'\\"forecast\\":\\"([^"\\]+)\\"', html)
    if not forecast_match:
        raise ValueError("Could not locate forecast key in rankings HTML")

    forecast_key = forecast_match.group(1)
    forecast_ts_match = re.search(r'\\"forecastFromTimestamp\\":(\d+)', html[forecast_match.start() : forecast_match.start() + 120])
    forecast_from_timestamp = int(forecast_ts_match.group(1)) if forecast_ts_match else None

    data_marker = '\\"data\\":'
    data_marker_idx = html.rfind(data_marker, 0, forecast_match.start())
    if data_marker_idx == -1:
        raise ValueError("Could not locate top chart data marker in rankings HTML")

    array_start = html.find("[", data_marker_idx)
    if array_start == -1:
        raise ValueError("Could not locate top chart data array in rankings HTML")

    raw_array = extract_balanced_block(html, array_start, "[", "]")
    decoded_array = bytes(raw_array, "utf-8").decode("unicode_escape")
    items = json.loads(decoded_array)

    rows: list[dict[str, object]] = []
    for item in items:
        row = {"week_start": item["x"]}
        row.update(item.get("ys", {}))
        rows.append(row)

    return {
        "forecast_key": forecast_key,
        "forecast_from_timestamp": forecast_from_timestamp,
        "rows": rows,
    }


def normalize_leaderboard_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    normalized: list[dict[str, object]] = []
    for row in rows:
        normalized.append(
            {
                **row,
                "weekly_tokens_billions": tokens_to_billions(str(row["weekly_tokens_display"])),
                "wow_change_ratio": float(row["wow_change_percent"]) / 100,
            }
        )
    return normalized


def enrich_leaderboard_rows(
    rows: list[dict[str, object]],
    generated_at: str,
    snapshot_captured_at: str | None = None,
) -> list[dict[str, object]]:
    snapshot_captured_at = snapshot_captured_at or snapshot_date_from_generated_at(generated_at)
    enriched: list[dict[str, object]] = []
    for row in rows:
        enriched.append(
            {
                **row,
                "source_url": row.get("source_url", SOURCE_URL),
                "ranking_period": row.get("ranking_period", RANKING_PERIOD),
                "snapshot_captured_at": row.get("snapshot_captured_at", snapshot_captured_at),
                "generated_at": row.get("generated_at", generated_at),
            }
        )
    return enriched


def ordered_timeseries_keys(timeseries_payload: dict[str, object]) -> list[str]:
    rows = timeseries_payload["rows"]
    seen = set()
    ordered: list[str] = []
    available = {key for row in rows for key in row.keys() if key != "week_start"}

    for series in TOP_CHART_SERIES_SNAPSHOT:
        key = series["dataKey"]
        if key in available and key not in seen:
            ordered.append(key)
            seen.add(key)

    for row in rows:
        for key in row.keys():
            if key != "week_start" and key not in seen:
                ordered.append(key)
                seen.add(key)

    return ordered


def merge_timeseries_history(existing_payload: dict[str, object], new_payload: dict[str, object]) -> dict[str, object]:
    merged_by_week = {
        str(row["week_start"]): dict(row)
        for row in existing_payload.get("rows", [])
        if row.get("week_start")
    }
    for row in new_payload.get("rows", []):
        if row.get("week_start"):
            merged_by_week[str(row["week_start"])] = dict(row)

    merged_rows = [merged_by_week[week] for week in sorted(merged_by_week)]
    return {
        "forecast_key": new_payload.get("forecast_key") or existing_payload.get("forecast_key", ""),
        "forecast_from_timestamp": new_payload.get("forecast_from_timestamp") or existing_payload.get("forecast_from_timestamp"),
        "rows": merged_rows,
    }


def leaderboard_history_key(row: dict[str, object]) -> tuple[str, str, str, str]:
    return (
        str(row.get("snapshot_captured_at", "")),
        str(row.get("ranking_period", RANKING_PERIOD)),
        str(row.get("rank", "")),
        str(row.get("model_url", row.get("model", ""))),
    )


def merge_leaderboard_history(existing_rows: list[dict[str, object]], new_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    merged = {leaderboard_history_key(row): dict(row) for row in existing_rows}
    for row in new_rows:
        merged[leaderboard_history_key(row)] = dict(row)
    return sorted(
        merged.values(),
        key=lambda row: (
            str(row.get("snapshot_captured_at", "")),
            int(row.get("rank", 0) or 0),
            str(row.get("model", "")),
        ),
    )


def latest_leaderboard_snapshot_date(rows: list[dict[str, object]]) -> str:
    snapshot_dates = [str(row.get("snapshot_captured_at", "")) for row in rows if row.get("snapshot_captured_at")]
    if snapshot_dates:
        return max(snapshot_dates)
    return LEADERBOARD_SNAPSHOT_CAPTURED_AT


def latest_leaderboard_snapshot_rows(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    latest_snapshot = latest_leaderboard_snapshot_date(rows)
    latest_rows = [dict(row) for row in rows if str(row.get("snapshot_captured_at", "")) == latest_snapshot]
    return sorted(latest_rows, key=lambda row: int(row.get("rank", 0) or 0))


def write_leaderboard_csv(rows: list[dict[str, object]], path: Path, generated_at: str) -> None:
    fieldnames = [
        "rank",
        "model",
        "provider",
        "weekly_tokens_display",
        "weekly_tokens_billions",
        "wow_change_percent",
        "wow_change_direction",
        "model_url",
        "provider_url",
        "source_url",
        "ranking_period",
        "snapshot_captured_at",
        "generated_at",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    "rank": row["rank"],
                    "model": row["model"],
                    "provider": row["provider"],
                    "weekly_tokens_display": row["weekly_tokens_display"],
                    "weekly_tokens_billions": row["weekly_tokens_billions"],
                    "wow_change_percent": row["wow_change_percent"],
                    "wow_change_direction": row["wow_change_direction"],
                    "model_url": row["model_url"],
                    "provider_url": row["provider_url"],
                    "source_url": row.get("source_url", SOURCE_URL),
                    "ranking_period": row.get("ranking_period", RANKING_PERIOD),
                    "snapshot_captured_at": row.get("snapshot_captured_at", snapshot_date_from_generated_at(generated_at)),
                    "generated_at": row.get("generated_at", generated_at),
                }
            )


def write_timeseries_csv(timeseries_payload: dict[str, object], path: Path) -> None:
    keys = ordered_timeseries_keys(timeseries_payload)
    fieldnames = ["week_start", *keys]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in timeseries_payload["rows"]:
            writer.writerow({key: row.get(key, "") for key in fieldnames})


def series_color_map() -> dict[str, str]:
    return {item["dataKey"]: item["fill"].lstrip("#") for item in TOP_CHART_SERIES_SNAPSHOT}


def clamp_channel(value: float) -> int:
    return max(0, min(255, int(round(value))))


def adjust_hex_color(hex_color: str, factor: float) -> str:
    channels = [int(hex_color[idx : idx + 2], 16) for idx in (0, 2, 4)]
    adjusted: list[int] = []
    for channel in channels:
        if factor >= 0:
            value = channel + (255 - channel) * factor
        else:
            value = channel * (1 + factor)
        adjusted.append(clamp_channel(value))
    return "".join(f"{channel:02X}" for channel in adjusted)


def extend_palette(base_palette: list[str], minimum_size: int) -> list[str]:
    extended = list(base_palette)
    if minimum_size <= len(extended):
        return extended

    index = 0
    while len(extended) < minimum_size:
        base_color = base_palette[index % len(base_palette)]
        cycle = index // len(base_palette)
        factor = min(0.8, 0.08 * (cycle + 1))
        adjustment = factor if cycle % 2 == 0 else -factor
        candidate = adjust_hex_color(base_color, adjustment)
        while candidate in extended and factor < 0.95:
            factor = min(0.95, factor + 0.03)
            adjustment = factor if cycle % 2 == 0 else -factor
            candidate = adjust_hex_color(base_color, adjustment)
        extended.append(candidate)
        index += 1
    return extended


def default_chart_series_keys(timeseries_payload: dict[str, object]) -> list[str]:
    forecast_key = timeseries_payload.get("forecast_key")
    return [key for key in ordered_timeseries_keys(timeseries_payload) if key != forecast_key]


def kimi_emphasis_series_color_map(series_keys: list[str]) -> dict[str, str]:
    color_map: dict[str, str] = {}
    if KIMI_SERIES_KEY in series_keys:
        color_map[KIMI_SERIES_KEY] = KIMI_HIGHLIGHT_COLOR
    if "Others" in series_keys:
        color_map["Others"] = KIMI_MUTED_COLOR

    remaining_keys = [key for key in series_keys if key not in {KIMI_SERIES_KEY, "Others"}]
    remaining_palette = [
        color
        for color in KIMI_REFERENCE_PALETTE
        if color not in {KIMI_HIGHLIGHT_COLOR, KIMI_MUTED_COLOR}
    ]
    extended_palette = extend_palette(remaining_palette, len(remaining_keys))
    for key, color in zip(remaining_keys, extended_palette):
        color_map[key] = color
    return color_map


def apply_series_colors(chart: BarChart | LineChart, series_keys: list[str], color_map: dict[str, str]) -> None:
    for idx, series in enumerate(chart.series):
        key = series_keys[idx]
        color = color_map[key]
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color


def normalize_series_lookup_key(series_key: str) -> str:
    if series_key == "Others":
        return series_key
    if "/" in series_key:
        provider, model = series_key.split("/", 1)
    else:
        provider, model = "", series_key
    model = model.split(":", 1)[0]
    while re.search(r"-\d{4,8}$", model):
        model = re.sub(r"-\d{4,8}$", "", model)
    return f"{provider}/{model}" if provider else model


def leaderboard_series_label_map(leaderboard_rows: list[dict[str, object]]) -> dict[str, str]:
    label_map: dict[str, str] = {}
    for row in leaderboard_rows:
        model_url = str(row.get("model_url", "")).rstrip("/")
        parts = model_url.split("/")
        if len(parts) >= 2:
            key = normalize_series_lookup_key("/".join(parts[-2:]))
            label_map[key] = str(row["model"])
    return label_map


def fallback_series_label(series_key: str) -> str:
    if series_key == "Others":
        return "Others"
    if series_key == KIMI_SERIES_KEY:
        return "Kimi K2.5"
    parts = series_key.split("/", 1)
    provider = parts[0] if len(parts) == 2 else ""
    label = parts[-1].replace(":free", " free")
    while re.search(r"-\d{4,8}$", label):
        label = re.sub(r"-\d{4,8}$", "", label)
    if provider and len(label) <= 3:
        label = f"{provider} {label}"
    label = label.replace("-", " ").replace("_", " ")
    return " ".join(label.split()).title()


def display_label_for_series_key(series_key: str, leaderboard_rows: list[dict[str, object]]) -> str:
    label_map = leaderboard_series_label_map(leaderboard_rows)
    return label_map.get(normalize_series_lookup_key(series_key), fallback_series_label(series_key))


def compact_chart_label(label: str) -> str:
    compact = label.replace(" (free)", " free")
    compact = compact.replace(" Plus ", " ")
    compact = compact.replace(" Preview Free", " Preview")
    compact = compact.replace(" Preview free", " Preview")
    compact = compact.replace(" K2.5", " K2.5")
    if " Preview" in compact and "\n" not in compact:
        compact = compact.replace(" Preview", "\nPreview", 1)
    elif compact.endswith(" free"):
        compact = compact[:-5] + "\nfree"
    elif len(compact) > 18 and " " in compact and "\n" not in compact:
        head, tail = compact.rsplit(" ", 1)
        compact = f"{head}\n{tail}"
    return compact


def select_display_trend_series_keys(timeseries_payload: dict[str, object], max_models: int = DISPLAY_TREND_MAX_MODELS) -> list[str]:
    rows = timeseries_payload["rows"]
    if not rows:
        return []

    latest_row = rows[-1]
    chart_series_keys = default_chart_series_keys(timeseries_payload)
    ranked_keys = sorted(
        [key for key in chart_series_keys if key != "Others"],
        key=lambda key: float(latest_row.get(key, 0) or 0),
        reverse=True,
    )
    selected = ranked_keys[:max_models]
    if KIMI_SERIES_KEY in chart_series_keys and KIMI_SERIES_KEY not in selected:
        selected.append(KIMI_SERIES_KEY)
    if "Others" in chart_series_keys:
        selected.append("Others")
    return selected


def prepare_display_datasets(timeseries_payload: dict[str, object], leaderboard_rows: list[dict[str, object]]) -> dict[str, object]:
    trend_series_keys = select_display_trend_series_keys(timeseries_payload)
    trend_color_map = kimi_emphasis_series_color_map(trend_series_keys)
    trend_headers = ["week_start", *[compact_chart_label(display_label_for_series_key(key, leaderboard_rows)) for key in trend_series_keys]]
    trend_rows = [
        [row["week_start"], *[row.get(key, 0) or 0 for key in trend_series_keys]]
        for row in timeseries_payload["rows"]
    ]
    legend_keys = [key for key in trend_series_keys if key not in {"Others", KIMI_SERIES_KEY}]
    if KIMI_SERIES_KEY in trend_series_keys:
        legend_keys.append(KIMI_SERIES_KEY)
    legend_entries = [
        {
            "key": key,
            "label": display_label_for_series_key(key, leaderboard_rows),
            "color": trend_color_map[key],
        }
        for key in legend_keys
    ]

    latest_row = timeseries_payload["rows"][-1] if timeseries_payload["rows"] else {}
    trend_summary = [
        {
            "color": trend_color_map[key],
            "model": display_label_for_series_key(key, leaderboard_rows),
            "value": format_token_amount(latest_row.get(key, 0)),
        }
        for key in trend_series_keys
    ]

    leaderboard_display_rows = leaderboard_rows[:8]
    leaderboard_sheet_rows = [[row["model"], row["weekly_tokens_billions"]] for row in leaderboard_display_rows]
    kimi_rows = [[row["week_start"], row.get(KIMI_SERIES_KEY, 0) or 0] for row in timeseries_payload["rows"]]
    kimi_values = [row[1] for row in kimi_rows]
    latest_kimi = kimi_values[-1] if kimi_values else 0
    total_latest = sum(float(latest_row.get(key, 0) or 0) for key in default_chart_series_keys(timeseries_payload))
    latest_share = (latest_kimi / total_latest * 100) if total_latest else 0
    kimi_summary = [
        {"label": "\u6700\u65b0\u5468", "value": format_token_amount(latest_kimi)},
        {"label": "\u5cf0\u503c", "value": format_token_amount(max(kimi_values) if kimi_values else 0)},
        {"label": "\u5468\u6570", "value": str(len(kimi_rows))},
        {"label": "\u6700\u65b0\u5360\u6bd4", "value": format_percent(latest_share)},
    ]
    wow_rows = leaderboard_display_rows
    wow_sheet_rows = [[row["model"], row["wow_change_percent"]] for row in wow_rows]

    return {
        "trend_series_keys": trend_series_keys,
        "trend_headers": trend_headers,
        "trend_rows": trend_rows,
        "trend_color_map": trend_color_map,
        "legend_entries": legend_entries,
        "trend_summary": trend_summary,
        "leaderboard_rows": leaderboard_display_rows,
        "leaderboard_sheet_rows": leaderboard_sheet_rows,
        "kimi_rows": kimi_rows,
        "kimi_summary": kimi_summary,
        "wow_rows": wow_rows,
        "wow_sheet_rows": wow_sheet_rows,
        "latest_week": latest_row.get("week_start", ""),
        "latest_snapshot_captured_at": latest_leaderboard_snapshot_date(leaderboard_rows),
    }


def clear_sheet(ws) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)
    ws._charts = []


def create_workbook_skeleton() -> Workbook:
    wb = Workbook()
    display_ws = wb.active
    display_ws.title = DISPLAY_SHEET_TITLE
    wb.create_sheet(META_SHEET)
    wb.create_sheet(RAW_TIMESERIES_SHEET)
    wb.create_sheet(RAW_LEADERBOARD_SHEET)
    wb.create_sheet(DATA_TREND_SHEET)
    wb.create_sheet(DATA_LEADERBOARD_SHEET)
    wb.create_sheet(DATA_KIMI_SHEET)
    wb.create_sheet(DATA_WOW_SHEET)
    wb.active = wb.sheetnames.index(DISPLAY_SHEET_TITLE)
    return wb


def ensure_workbook_skeleton(wb: Workbook) -> Workbook:
    required_titles = [
        DISPLAY_SHEET_TITLE,
        META_SHEET,
        RAW_TIMESERIES_SHEET,
        RAW_LEADERBOARD_SHEET,
        DATA_TREND_SHEET,
        DATA_LEADERBOARD_SHEET,
        DATA_KIMI_SHEET,
        DATA_WOW_SHEET,
    ]
    for title in required_titles:
        if title not in wb.sheetnames:
            wb.create_sheet(title)
    if "Charts" in wb.sheetnames:
        wb.remove(wb["Charts"])
    wb.active = wb.sheetnames.index(DISPLAY_SHEET_TITLE)
    return wb


def read_meta_values(wb: Workbook) -> dict[str, str]:
    if META_SHEET not in wb.sheetnames:
        return {}
    ws = wb[META_SHEET]
    values: dict[str, str] = {}
    for key, value in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        if key:
            values[str(key)] = "" if value is None else str(value)
    return values


def write_meta_sheet(
    ws,
    generated_at: str,
    timeseries_payload: dict[str, object],
    leaderboard_rows: list[dict[str, object]],
    data_refreshed_at: str,
    display_refreshed_at: str,
) -> None:
    clear_sheet(ws)
    title_fill = PatternFill("solid", fgColor="0F172A")
    note_fill = PatternFill("solid", fgColor="E2E8F0")
    text_light = "F8FAFC"

    ws["A1"] = "OpenRouter Rankings Snapshot"
    ws["A1"].font = Font(size=16, bold=True, color=text_light)
    ws["A1"].fill = title_fill

    meta_rows = [
        ("source_url", SOURCE_URL),
        ("generated_at", generated_at),
        ("ranking_period", RANKING_PERIOD),
        ("data_refreshed_at", data_refreshed_at),
        ("display_refreshed_at", display_refreshed_at),
        ("top_chart_extraction_method", "Directly parsed from embedded rankings HTML data block."),
        ("leaderboard_extraction_method", "Direct DOM snapshot extracted after expanding Show more on 2026-03-31."),
        ("leaderboard_snapshot_captured_at", latest_leaderboard_snapshot_date(leaderboard_rows)),
        ("timeseries_week_count", len(timeseries_payload.get("rows", []))),
        ("leaderboard_snapshot_count", len({str(row.get("snapshot_captured_at", "")) for row in leaderboard_rows if row.get("snapshot_captured_at")})),
        ("forecast_key", timeseries_payload.get("forecast_key", "")),
        ("palette_source", KIMI_PALETTE_SOURCE_NOTE),
        ("notes", "Single workbook with raw data layer plus a display dashboard. No values are estimated."),
    ]
    for row_idx, (label, value) in enumerate(meta_rows, start=2):
        ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row_idx, column=2, value=value)
        ws.cell(row=row_idx, column=1).fill = note_fill
        ws.cell(row=row_idx, column=2).fill = note_fill
        ws.cell(row=row_idx, column=2).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 96


def write_table_sheet(ws, headers: list[str], rows: list[list[object]], number_formats: dict[int, str] | None = None) -> None:
    clear_sheet(ws)
    header_fill = PatternFill("solid", fgColor="1E293B")
    text_light = "F8FAFC"

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color=text_light)
        cell.fill = header_fill

    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
            if number_formats and col_idx in number_formats:
                ws.cell(row=row_idx, column=col_idx).number_format = number_formats[col_idx]

    ws.freeze_panes = "A2"
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18


def write_raw_time_sheet(ws, timeseries_payload: dict[str, object]) -> None:
    keys = ordered_timeseries_keys(timeseries_payload)
    rows = [[row["week_start"], *[row.get(key) for key in keys]] for row in timeseries_payload["rows"]]
    write_table_sheet(ws, ["week_start", *keys], rows)
    ws.column_dimensions["A"].width = 14


def write_raw_leaderboard_sheet(ws, leaderboard_rows: list[dict[str, object]]) -> None:
    headers = [
        "rank",
        "model",
        "provider",
        "weekly_tokens_display",
        "weekly_tokens_billions",
        "wow_change_percent",
        "wow_change_direction",
        "model_url",
        "provider_url",
        "source_url",
        "ranking_period",
        "snapshot_captured_at",
        "generated_at",
    ]
    rows = [
        [
            row["rank"],
            row["model"],
            row["provider"],
            row["weekly_tokens_display"],
            row["weekly_tokens_billions"],
            row["wow_change_percent"],
            row["wow_change_direction"],
            row["model_url"],
            row["provider_url"],
            row.get("source_url", SOURCE_URL),
            row.get("ranking_period", RANKING_PERIOD),
            row.get("snapshot_captured_at", ""),
            row.get("generated_at", ""),
        ]
        for row in leaderboard_rows
    ]
    write_table_sheet(ws, headers, rows, number_formats={5: '0.00 "B"', 6: '0;[Red]-0'})
    widths = {"A": 8, "B": 28, "C": 16, "D": 18, "E": 18, "F": 18, "G": 16, "H": 42, "I": 28, "J": 28, "K": 16, "L": 20, "M": 24}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def write_display_helper_sheets(wb: Workbook, display_data: dict[str, object]) -> None:
    write_table_sheet(wb[DATA_TREND_SHEET], display_data["trend_headers"], display_data["trend_rows"])
    write_table_sheet(
        wb[DATA_LEADERBOARD_SHEET],
        ["model", "weekly_tokens_billions"],
        display_data["leaderboard_sheet_rows"],
        number_formats={2: '0.00 "B"'},
    )
    write_table_sheet(wb[DATA_KIMI_SHEET], ["week_start", "kimi_usage"], display_data["kimi_rows"])
    write_table_sheet(
        wb[DATA_WOW_SHEET],
        ["model", "wow_change_percent"],
        display_data["wow_sheet_rows"],
        number_formats={2: '0;[Red]-0'},
    )
    for title in (DATA_TREND_SHEET, DATA_LEADERBOARD_SHEET, DATA_KIMI_SHEET, DATA_WOW_SHEET):
        wb[title].sheet_state = "hidden"


def write_display_placeholder(ws, generated_at: str) -> None:
    clear_sheet(ws)
    ws.sheet_view.showGridLines = False
    ws["A1"] = DISPLAY_PAGE_TITLE
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = DISPLAY_PLACEHOLDER_NOTE
    ws["A3"] = f"{DISPLAY_DATA_REFRESHED_LABEL}{generated_at}"
    configure_display_page_setup(ws)


def write_display_summary_headers(ws, start_row: int, start_col: int, headers: list[str]) -> None:
    header_fill = PatternFill("solid", fgColor="E8EEF7")
    for offset, header in enumerate(headers):
        cell = ws.cell(row=start_row, column=start_col + offset, value=header)
        cell.font = Font(bold=True, size=9)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def style_display_summary_range(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    for row_idx in range(start_row, end_row + 1):
        ws.row_dimensions[row_idx].height = 15
        for col_idx in range(start_col, end_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical="center", shrink_to_fit=True)


def build_display_sheet(ws, display_data: dict[str, object], generated_at: str) -> None:
    clear_sheet(ws)
    ws.sheet_view.showGridLines = False

    ws["A1"] = DISPLAY_PAGE_TITLE
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = DISPLAY_PAGE_SUBTITLE
    ws["A3"] = f"{DISPLAY_REFRESHED_AT_LABEL}{generated_at}"
    ws["B4"] = DISPLAY_TOP_LEFT_TITLE
    ws["M4"] = DISPLAY_TOP_RIGHT_TITLE
    ws["B30"] = DISPLAY_BOTTOM_LEFT_TITLE
    ws["M30"] = DISPLAY_BOTTOM_RIGHT_TITLE
    for cell_ref in ("B4", "M4", "B30", "M30"):
        ws[cell_ref].font = Font(bold=True, size=11)

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 9
    ws.column_dimensions["C"].width = 37
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["M"].width = 9
    ws.column_dimensions["N"].width = 37
    ws.column_dimensions["O"].width = 18
    for row_idx, height in {1: 24, 2: 18, 3: 16, 4: 18, 5: 8, 19: 18, 20: 16, 30: 18, 45: 18, 46: 16}.items():
        ws.row_dimensions[row_idx].height = height

    latest_week = display_data["latest_week"]
    ws["B19"] = f"{DISPLAY_LATEST_WEEK_LABEL}{latest_week}"
    ws["M19"] = f"{DISPLAY_SNAPSHOT_DATE_LABEL}{display_data['latest_snapshot_captured_at']}"
    ws["B45"] = DISPLAY_KIMI_METRICS_TITLE
    ws["M45"] = DISPLAY_WOW_SUMMARY_TITLE
    for cell_ref in ("B19", "M19", "B45", "M45"):
        ws[cell_ref].font = Font(bold=True, size=11)

    write_display_summary_headers(ws, 20, 2, ["\u989c\u8272", "\u6a21\u578b", "\u6700\u65b0\u5468\u7528\u91cf"])
    for row_idx, summary in enumerate(display_data["trend_summary"][:8], start=21):
        color_cell = ws.cell(row=row_idx, column=2, value=" ")
        color_cell.fill = PatternFill("solid", fgColor=summary["color"])
        ws.cell(row=row_idx, column=3, value=summary["model"])
        ws.cell(row=row_idx, column=4, value=summary["value"])
    style_display_summary_range(ws, 21, 28, 2, 4)

    write_display_summary_headers(ws, 20, 13, ["\u6392\u540d", "\u6a21\u578b", "\u672c\u5468\u7528\u91cf"])
    for row_idx, row in enumerate(display_data["leaderboard_rows"], start=21):
        ws.cell(row=row_idx, column=13, value=row["rank"])
        ws.cell(row=row_idx, column=14, value=row["model"])
        ws.cell(row=row_idx, column=15, value=row["weekly_tokens_display"])
    style_display_summary_range(ws, 21, 28, 13, 15)

    write_display_summary_headers(ws, 46, 2, ["\u6307\u6807", "\u6570\u503c"])
    for row_idx, row in enumerate(display_data["kimi_summary"], start=47):
        ws.cell(row=row_idx, column=2, value=row["label"])
        ws.cell(row=row_idx, column=3, value=row["value"])
    style_display_summary_range(ws, 47, 50, 2, 3)

    write_display_summary_headers(ws, 46, 13, ["\u6392\u540d", "\u6a21\u578b", "WoW"])
    for row_idx, row in enumerate(display_data["wow_rows"], start=47):
        ws.cell(row=row_idx, column=13, value=row["rank"])
        ws.cell(row=row_idx, column=14, value=row["model"])
        ws.cell(row=row_idx, column=15, value=f"{row['wow_change_percent']}%")
    style_display_summary_range(ws, 47, 54, 13, 15)

    trend_data_max_row = ws.parent[DATA_TREND_SHEET].max_row

    trend_chart = BarChart()
    trend_chart.type = "col"
    trend_chart.grouping = "stacked"
    trend_chart.style = 13
    trend_chart.y_axis.title = "Tokens"
    trend_chart.legend.position = "r"
    trend_chart.legend.overlay = False
    trend_chart.width = 13.6
    trend_chart.height = 6.0
    trend_chart.gapWidth = 0
    trend_chart.overlap = 100
    trend_chart.add_data(
        Reference(ws.parent[DATA_TREND_SHEET], min_col=2, max_col=len(display_data["trend_headers"]), min_row=1, max_row=trend_data_max_row),
        titles_from_data=True,
    )
    trend_chart.set_categories(
        Reference(ws.parent[DATA_TREND_SHEET], min_col=1, min_row=2, max_row=trend_data_max_row)
    )
    apply_series_colors(trend_chart, display_data["trend_series_keys"], display_data["trend_color_map"])
    ws.add_chart(trend_chart, "B6")

    leaderboard_chart = BarChart()
    leaderboard_chart.type = "bar"
    leaderboard_chart.style = 10
    leaderboard_chart.x_axis.title = "Weekly Tokens (Billions)"
    leaderboard_chart.legend = None
    leaderboard_chart.width = 13.6
    leaderboard_chart.height = 6.0
    leaderboard_chart.gapWidth = 35
    leaderboard_chart.add_data(
        Reference(ws.parent[DATA_LEADERBOARD_SHEET], min_col=2, min_row=1, max_row=max(ws.parent[DATA_LEADERBOARD_SHEET].max_row, MAX_LEADERBOARD_ROWS)),
        titles_from_data=True,
    )
    leaderboard_chart.set_categories(
        Reference(ws.parent[DATA_LEADERBOARD_SHEET], min_col=1, min_row=2, max_row=max(ws.parent[DATA_LEADERBOARD_SHEET].max_row, MAX_LEADERBOARD_ROWS))
    )
    leaderboard_chart.y_axis.reverseOrder = True
    if leaderboard_chart.series:
        leaderboard_chart.series[0].graphicalProperties.solidFill = KIMI_LEADERBOARD_COLOR
        leaderboard_chart.series[0].graphicalProperties.line.solidFill = KIMI_LEADERBOARD_COLOR
    ws.add_chart(leaderboard_chart, "M6")

    kimi_chart = LineChart()
    kimi_chart.style = 13
    kimi_chart.y_axis.title = "Tokens"
    kimi_chart.legend = None
    kimi_chart.width = 13.6
    kimi_chart.height = 6.0
    kimi_chart.add_data(
        Reference(ws.parent[DATA_KIMI_SHEET], min_col=2, min_row=1, max_row=max(ws.parent[DATA_KIMI_SHEET].max_row, MAX_KIMI_ROWS)),
        titles_from_data=True,
    )
    kimi_chart.set_categories(
        Reference(ws.parent[DATA_KIMI_SHEET], min_col=1, min_row=2, max_row=max(ws.parent[DATA_KIMI_SHEET].max_row, MAX_KIMI_ROWS))
    )
    kimi_chart.series[0].marker.symbol = "circle"
    kimi_chart.series[0].marker.size = 7
    kimi_chart.series[0].graphicalProperties.line.solidFill = KIMI_HIGHLIGHT_COLOR
    kimi_chart.series[0].graphicalProperties.solidFill = KIMI_HIGHLIGHT_COLOR
    ws.add_chart(kimi_chart, "B32")

    wow_chart = BarChart()
    wow_chart.type = "bar"
    wow_chart.style = 11
    wow_chart.x_axis.title = "WoW Change (%)"
    wow_chart.legend = None
    wow_chart.width = 13.6
    wow_chart.height = 6.0
    wow_chart.gapWidth = 35
    wow_chart.add_data(
        Reference(ws.parent[DATA_WOW_SHEET], min_col=2, min_row=1, max_row=max(ws.parent[DATA_WOW_SHEET].max_row, MAX_WOW_ROWS)),
        titles_from_data=True,
    )
    wow_chart.set_categories(
        Reference(ws.parent[DATA_WOW_SHEET], min_col=1, min_row=2, max_row=max(ws.parent[DATA_WOW_SHEET].max_row, MAX_WOW_ROWS))
    )
    wow_chart.y_axis.reverseOrder = True
    if wow_chart.series:
        wow_chart.series[0].graphicalProperties.solidFill = KIMI_WOW_COLOR
        wow_chart.series[0].graphicalProperties.line.solidFill = KIMI_WOW_COLOR
    ws.add_chart(wow_chart, "M32")
    configure_display_page_setup(ws)


def configure_display_page_setup(ws) -> None:
    ws.print_area = "A1:Q55"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.2
    ws.page_margins.bottom = 0.2
    ws.page_margins.header = 0.1
    ws.page_margins.footer = 0.1


def load_timeseries_payload_from_workbook(wb: Workbook) -> dict[str, object]:
    ws = wb[RAW_TIMESERIES_SHEET]
    headers = [cell.value for cell in ws[1]]
    if not any(headers):
        return {"forecast_key": "", "rows": []}
    rows: list[dict[str, object]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        row_dict = {
            str(headers[idx]): value
            for idx, value in enumerate(row)
            if idx < len(headers) and headers[idx] is not None and value is not None
        }
        rows.append(row_dict)

    meta_values = read_meta_values(wb)
    forecast_key = meta_values.get("forecast_key", "")
    if not forecast_key:
        for header in headers:
            if header and str(header).startswith("forecast-"):
                forecast_key = str(header)
                break
    return {"forecast_key": forecast_key, "rows": rows}


def load_leaderboard_rows_from_workbook(wb: Workbook, latest_snapshot_only: bool = False) -> list[dict[str, object]]:
    ws = wb[RAW_LEADERBOARD_SHEET]
    headers = [cell.value for cell in ws[1]]
    if not any(headers):
        return []
    rows: list[dict[str, object]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        row_dict = {
            str(headers[idx]): value
            for idx, value in enumerate(row)
            if idx < len(headers) and headers[idx] is not None
        }
        row_dict["wow_change_ratio"] = float(row_dict["wow_change_percent"]) / 100
        if not row_dict.get("source_url"):
            row_dict["source_url"] = SOURCE_URL
        if not row_dict.get("ranking_period"):
            row_dict["ranking_period"] = RANKING_PERIOD
        if not row_dict.get("snapshot_captured_at"):
            row_dict["snapshot_captured_at"] = read_meta_values(wb).get("leaderboard_snapshot_captured_at", LEADERBOARD_SNAPSHOT_CAPTURED_AT)
        rows.append(row_dict)
    if latest_snapshot_only:
        return latest_leaderboard_snapshot_rows(rows)
    return merge_leaderboard_history([], rows)


def remove_deprecated_variant_workbook(workbook_path: Path) -> None:
    deprecated_path = workbook_path.parent / "openrouter_top_models_kimi_emphasis.xlsx"
    if deprecated_path.exists():
        deprecated_path.unlink()


def write_data_layers(
    wb: Workbook,
    timeseries_payload: dict[str, object],
    leaderboard_rows: list[dict[str, object]],
    generated_at: str,
    display_refreshed_at: str,
) -> tuple[dict[str, object], list[dict[str, object]]]:
    ensure_workbook_skeleton(wb)
    existing_timeseries_payload = load_timeseries_payload_from_workbook(wb)
    existing_leaderboard_rows = load_leaderboard_rows_from_workbook(wb, latest_snapshot_only=False)
    merged_timeseries_payload = merge_timeseries_history(existing_timeseries_payload, timeseries_payload)
    merged_leaderboard_rows = merge_leaderboard_history(existing_leaderboard_rows, leaderboard_rows)
    latest_leaderboard_rows = latest_leaderboard_snapshot_rows(merged_leaderboard_rows)
    display_data = prepare_display_datasets(merged_timeseries_payload, latest_leaderboard_rows)
    write_meta_sheet(
        wb[META_SHEET],
        generated_at=generated_at,
        timeseries_payload=merged_timeseries_payload,
        leaderboard_rows=merged_leaderboard_rows,
        data_refreshed_at=generated_at,
        display_refreshed_at=display_refreshed_at,
    )
    write_raw_time_sheet(wb[RAW_TIMESERIES_SHEET], merged_timeseries_payload)
    write_raw_leaderboard_sheet(wb[RAW_LEADERBOARD_SHEET], merged_leaderboard_rows)
    write_display_helper_sheets(wb, display_data)
    return merged_timeseries_payload, merged_leaderboard_rows


def refresh_display_layer(wb: Workbook, generated_at: str) -> None:
    ensure_workbook_skeleton(wb)
    timeseries_payload = load_timeseries_payload_from_workbook(wb)
    leaderboard_history_rows = load_leaderboard_rows_from_workbook(wb, latest_snapshot_only=False)
    leaderboard_rows = latest_leaderboard_snapshot_rows(leaderboard_history_rows)
    display_data = prepare_display_datasets(timeseries_payload, leaderboard_rows)
    meta_values = read_meta_values(wb)
    data_refreshed_at = meta_values.get("data_refreshed_at", generated_at)
    write_meta_sheet(
        wb[META_SHEET],
        generated_at=generated_at,
        timeseries_payload=timeseries_payload,
        leaderboard_rows=leaderboard_history_rows,
        data_refreshed_at=data_refreshed_at,
        display_refreshed_at=generated_at,
    )
    write_display_helper_sheets(wb, display_data)
    build_display_sheet(wb[DISPLAY_SHEET_TITLE], display_data, generated_at)


def build_workbook(
    timeseries_payload: dict[str, object],
    leaderboard_rows: list[dict[str, object]],
    workbook_path: Path,
    generated_at: str | None = None,
) -> None:
    generated_at = generated_at or current_timestamp()
    wb = create_workbook_skeleton()
    merged_timeseries_payload, merged_leaderboard_rows = write_data_layers(wb, timeseries_payload, leaderboard_rows, generated_at, generated_at)
    write_timeseries_csv(merged_timeseries_payload, workbook_path.parent / "openrouter_top_models_timeseries.csv")
    write_leaderboard_csv(merged_leaderboard_rows, workbook_path.parent / "openrouter_top_models.csv", generated_at)
    refresh_display_layer(wb, generated_at)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(workbook_path)


def run_data_refresh(
    workbook_path: Path,
    leaderboard_csv_path: Path,
    timeseries_csv_path: Path,
    cache_path: Path,
) -> None:
    generated_at = current_timestamp()
    html = load_rankings_html(cache_path)
    timeseries_payload = extract_top_models_chart_payload_from_html(html)
    leaderboard_rows = enrich_leaderboard_rows(normalize_leaderboard_rows(LEADERBOARD_SNAPSHOT), generated_at)

    if workbook_path.exists():
        wb = load_workbook(workbook_path)
        display_refreshed_at = read_meta_values(wb).get("display_refreshed_at", "")
        merged_timeseries_payload, merged_leaderboard_rows = write_data_layers(wb, timeseries_payload, leaderboard_rows, generated_at, display_refreshed_at)
    else:
        wb = create_workbook_skeleton()
        merged_timeseries_payload, merged_leaderboard_rows = write_data_layers(wb, timeseries_payload, leaderboard_rows, generated_at, "")
        write_display_placeholder(wb[DISPLAY_SHEET_TITLE], generated_at)

    write_leaderboard_csv(merged_leaderboard_rows, leaderboard_csv_path, generated_at)
    write_timeseries_csv(merged_timeseries_payload, timeseries_csv_path)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(workbook_path)
    remove_deprecated_variant_workbook(workbook_path)


def run_display_refresh(workbook_path: Path) -> None:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found for display-refresh: {workbook_path}")
    generated_at = current_timestamp()
    wb = load_workbook(workbook_path)
    refresh_display_layer(wb, generated_at)
    wb.save(workbook_path)
    remove_deprecated_variant_workbook(workbook_path)


def run_full(
    workbook_path: Path,
    leaderboard_csv_path: Path,
    timeseries_csv_path: Path,
    cache_path: Path,
) -> None:
    generated_at = current_timestamp()
    html = load_rankings_html(cache_path)
    timeseries_payload = extract_top_models_chart_payload_from_html(html)
    leaderboard_rows = enrich_leaderboard_rows(normalize_leaderboard_rows(LEADERBOARD_SNAPSHOT), generated_at)

    if workbook_path.exists():
        wb = load_workbook(workbook_path)
        display_refreshed_at = read_meta_values(wb).get("display_refreshed_at", "")
        merged_timeseries_payload, merged_leaderboard_rows = write_data_layers(
            wb,
            timeseries_payload,
            leaderboard_rows,
            generated_at,
            display_refreshed_at,
        )
        refresh_display_layer(wb, generated_at)
        workbook_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(workbook_path)
    else:
        build_workbook(timeseries_payload, leaderboard_rows, workbook_path, generated_at=generated_at)
        wb = load_workbook(workbook_path)
        merged_timeseries_payload = load_timeseries_payload_from_workbook(wb)
        merged_leaderboard_rows = load_leaderboard_rows_from_workbook(wb, latest_snapshot_only=False)

    write_leaderboard_csv(merged_leaderboard_rows, leaderboard_csv_path, generated_at)
    write_timeseries_csv(merged_timeseries_payload, timeseries_csv_path)
    remove_deprecated_variant_workbook(workbook_path)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Generate the OpenRouter rankings workbook.")
    parser.add_argument("--mode", choices=["full", "data-refresh", "display-refresh"], default="full")
    parser.add_argument("--workbook-path", type=Path)
    parser.add_argument("--leaderboard-csv-path", type=Path)
    parser.add_argument("--timeseries-csv-path", type=Path)
    parser.add_argument("--cache-path", type=Path)
    return parser


def main(argv: list[str] | None = None) -> None:
    parser = build_arg_parser()
    args = parser.parse_args(argv)

    root = Path.cwd()
    output_dir = root / "output" / "spreadsheet"
    tmp_dir = root / "tmp" / "spreadsheets"

    workbook_path = args.workbook_path or output_dir / "openrouter_top_models.xlsx"
    leaderboard_csv_path = args.leaderboard_csv_path or output_dir / "openrouter_top_models.csv"
    timeseries_csv_path = args.timeseries_csv_path or output_dir / "openrouter_top_models_timeseries.csv"
    cache_path = args.cache_path or tmp_dir / "rankings_page.html"

    if args.mode == "full":
        run_full(workbook_path, leaderboard_csv_path, timeseries_csv_path, cache_path)
    elif args.mode == "data-refresh":
        run_data_refresh(workbook_path, leaderboard_csv_path, timeseries_csv_path, cache_path)
    else:
        run_display_refresh(workbook_path)

    print(leaderboard_csv_path)
    print(timeseries_csv_path)
    print(workbook_path)


if __name__ == "__main__":
    main()
